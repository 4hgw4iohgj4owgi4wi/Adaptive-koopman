from __future__ import annotations

import csv
import shutil
from pathlib import Path

from docx import Document
from docx.shared import Inches


ROOT = Path(r"D:\LEARNING\ZNN\ZNN\Adaptive-koopman\Adaptive-koopman-main")
BASE = ROOT / "paper_dcn_tf12_draft" / "manuscript_zh_literature_update_figures_tf14_lyapunov_2026-05-10.docx"
OUT = ROOT / "paper_dcn_tf12_draft" / "manuscript_zh_tf14_round10_n20_e0e2_comm_update_2026-05-11.docx"

E0_SUMMARY = ROOT / "paper_dcn_tf12_draft" / "publication_readiness_round8_2026-05-11" / "e0_n20_mixed_fault_main_comparison" / "data" / "tf14_remaining_summary.csv"
E0_COMM_SUMMARY = ROOT / "paper_dcn_tf12_draft" / "publication_readiness_round10_2026-05-11" / "e0_n20_comm_noise_key_main_comparison_merged" / "data" / "tf14_remaining_summary.csv"
E2_SUMMARY = ROOT / "paper_dcn_tf12_draft" / "publication_readiness_round10_2026-05-11" / "e2_n20_mixed_fault_ablation_merged" / "data" / "tf14_remaining_summary.csv"
E2_COMM_SUMMARY = ROOT / "paper_dcn_tf12_draft" / "publication_readiness_round10_2026-05-11" / "e2_n20_comm_noise_key_ablation_merged" / "data" / "tf14_remaining_summary.csv"
E3_SUMMARY = ROOT / "paper_dcn_tf12_draft" / "publication_readiness_round8_2026-05-11" / "e3_n20_certificate" / "data" / "tf14_remaining_summary.csv"
E7_SUMMARY = ROOT / "paper_dcn_tf12_draft" / "publication_readiness_round8_2026-05-11" / "e7_n20_connection_safety" / "data" / "tf14_remaining_summary.csv"

FIGS = {
    "r0_summary": ROOT / "paper_dcn_tf12_draft" / "publication_readiness_round8_2026-05-11" / "e0_n20_mixed_fault_main_comparison" / "figures" / "R0_main_comparison_summary.png",
    "r0_timeseries": ROOT / "paper_dcn_tf12_draft" / "publication_readiness_round8_2026-05-11" / "e0_n20_mixed_fault_main_comparison" / "figures" / "R0_error_timeseries.png",
    "r0_comm_summary": ROOT / "paper_dcn_tf12_draft" / "publication_readiness_round10_2026-05-11" / "e0_n20_comm_noise_key_main_comparison_merged" / "figures" / "R0_comm_key_main_comparison_summary.png",
    "r0_comm_relative": ROOT / "paper_dcn_tf12_draft" / "publication_readiness_round10_2026-05-11" / "e0_n20_comm_noise_key_main_comparison_merged" / "figures" / "R0_comm_key_relative_change.png",
    "r2_heatmap": ROOT / "paper_dcn_tf12_draft" / "publication_readiness_round10_2026-05-11" / "e2_n20_mixed_fault_ablation_merged" / "figures" / "R2_round10_module_ablation_heatmap.png",
    "r2_effects": ROOT / "paper_dcn_tf12_draft" / "publication_readiness_round10_2026-05-11" / "e2_n20_mixed_fault_ablation_merged" / "figures" / "R2_round10_key_module_effects.png",
    "r2_comm_heatmap": ROOT / "paper_dcn_tf12_draft" / "publication_readiness_round10_2026-05-11" / "e2_n20_comm_noise_key_ablation_merged" / "figures" / "R2_comm_key_ablation_heatmap.png",
    "r2_comm_effects": ROOT / "paper_dcn_tf12_draft" / "publication_readiness_round10_2026-05-11" / "e2_n20_comm_noise_key_ablation_merged" / "figures" / "R2_comm_key_effects.png",
    "r3_cert": ROOT / "paper_dcn_tf12_draft" / "publication_readiness_round8_2026-05-11" / "e3_n20_certificate" / "figures" / "R3_certificate_statistics.png",
    "r7_comm": ROOT / "paper_dcn_tf12_draft" / "publication_readiness_round8_2026-05-11" / "e7_n20_connection_safety" / "figures" / "R7_payload_connection_safety_sine_comm_noise_high.png",
    "r7_mixed": ROOT / "paper_dcn_tf12_draft" / "publication_readiness_round8_2026-05-11" / "e7_n20_connection_safety" / "figures" / "R7_payload_connection_safety_sine_mixed_fault_noise.png",
}


def rows(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def f(row: dict[str, str], key: str) -> float:
    try:
        return float(row.get(key, "nan"))
    except Exception:
        return float("nan")


def fmt(x: float, digits: int = 4) -> str:
    if x != x:
        return "-"
    return f"{x:.{digits}f}"


def pct_delta(new: float, base: float) -> str:
    if base == 0 or base != base or new != new:
        return "-"
    return f"{(new - base) / base * 100:.2f}%"


def pct_reduction(new: float, base: float) -> str:
    if base == 0 or base != base or new != new:
        return "-"
    return f"{abs((new - base) / base * 100):.2f}%"


def short_scenario(name: str) -> str:
    return {
        "sine_comm_noise_high": "comm-high",
        "sine_mixed_fault_noise": "mixed",
        "sine_nominal_clean": "nominal",
        "sine_single_fault_v2": "single-fault",
    }.get(name, name)


def short_method(name: str) -> str:
    return {
        "baseline": "AKE-M",
        "tf14_phase_role": "TF14-full",
        "tf14_main": "TF14-main",
        "full_tf14": "TF14-full",
        "no_bilinear": "No bilinear",
        "no_comm_aware": "No comm-aware",
        "no_delay_compensation": "No delay comp.",
        "no_fdi": "No FDI",
        "no_ftc_switching": "No FTC",
        "no_online_adapt": "No online adapt",
        "no_phase_role": "No phase-role",
        "no_ppc_progress_guard": "No PPC/progress",
        "no_realtime_scheduler": "No RT scheduler",
        "no_stable_projection": "No stable proj.",
        "zoh_consensus_surrogate": "ZOH",
    }.get(name, name)


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph(text)
    try:
        p.style = "Caption"
    except Exception:
        pass


def add_section_heading(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    try:
        p.style = "Heading 2"
    except Exception:
        pass


def add_metric_table(doc: Document, title: str, headers: list[str], data: list[list[str]]) -> None:
    doc.add_paragraph(title)
    table = doc.add_table(rows=1, cols=len(headers))
    try:
        table.style = "Table Grid"
    except Exception:
        pass
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
    for row in data:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value


def remove_paragraph(paragraph) -> None:
    paragraph._element.getparent().remove(paragraph._element)


def cleanup_round10_heading(doc: Document) -> None:
    heading_text = "5.8 Round10 多随机种子结果与 claim 边界更新"
    content_prefix = "为将上一版代表性图包推进到可审稿的数据口径"
    content_para = next((p for p in doc.paragraphs if p.text.strip().startswith(content_prefix)), None)
    if content_para is None:
        return
    heading = doc.add_paragraph()
    run = heading.add_run(heading_text)
    run.bold = True
    content_para._p.addprevious(heading._p)

    for p in list(doc.paragraphs):
        if p.text.strip() != heading_text:
            continue
        next_p = p._element.getnext()
        next_text = ""
        if next_p is not None:
            try:
                next_text = "".join(next_p.itertext()).strip()
            except Exception:
                next_text = ""
        if not next_text.startswith(content_prefix):
            remove_paragraph(p)


def insert_new_section(doc: Document, before_text: str) -> None:
    # Build new content at the end, then move it before the conclusion heading.
    start = len(doc.element.body)

    add_section_heading(doc, "5.8 Round10 多随机种子结果与 claim 边界更新")
    doc.add_paragraph(
        "为将上一版代表性图包推进到可审稿的数据口径，本轮补充了 E0 targeted mixed fault/noise 主对比、E0 high-communication-noise 主对比、"
        "E2 mixed fault/noise 全量模块消融、E2 high-communication-noise 关键模块消融、E3 证书统计和 E7 载荷连接安全六组 paired n=20 结果。"
        "这些结果均由同一 runner 导出 run-level CSV、summary CSV、manifest 和 gate 表，并保留 method_id、baseline_class、trace_id、success、failure_type、"
        "connection_violation_count、force_norm_peak 等 provenance 字段。需要强调的是，当前 baseline 仍是 task-aligned AKE-M mechanism comparator，"
        "不是上传原文 AKE-A 的严格复现。"
    )

    e0 = rows(E0_SUMMARY)
    e0_by_method = {r["method"]: r for r in e0}
    base = e0_by_method["baseline"]
    full = e0_by_method["tf14_phase_role"]
    main = e0_by_method["tf14_main"]
    zoh = e0_by_method["zoh_consensus_surrogate"]
    add_metric_table(
        doc,
        "表X  E0 mixed fault/noise targeted 主对比（paired n=20）",
        ["方法", "路径完成率", "横向RMSE均值/m", "纵向RMSE均值/m", "连接最大利用率", "力峰值/N"],
        [
            ["AKE-M baseline", fmt(f(base, "full_path_reached_mean"), 2), fmt(f(base, "rmse_lat_mean_mean")), fmt(f(base, "rmse_long_mean_mean")), fmt(f(base, "connection_max_utilization_mean")), fmt(f(base, "force_norm_peak_mean"), 1)],
            ["TF14-main", fmt(f(main, "full_path_reached_mean"), 2), fmt(f(main, "rmse_lat_mean_mean")), fmt(f(main, "rmse_long_mean_mean")), fmt(f(main, "connection_max_utilization_mean")), fmt(f(main, "force_norm_peak_mean"), 1)],
            ["TF14-full", fmt(f(full, "full_path_reached_mean"), 2), fmt(f(full, "rmse_lat_mean_mean")), fmt(f(full, "rmse_long_mean_mean")), fmt(f(full, "connection_max_utilization_mean")), fmt(f(full, "force_norm_peak_mean"), 1)],
            ["ZOH surrogate", fmt(f(zoh, "full_path_reached_mean"), 2), fmt(f(zoh, "rmse_lat_mean_mean")), fmt(f(zoh, "rmse_long_mean_mean")), fmt(f(zoh, "connection_max_utilization_mean")), fmt(f(zoh, "force_norm_peak_mean"), 1)],
        ],
    )
    doc.add_paragraph(
        "在 mixed fault/noise targeted 主对比中，TF14-full 相对 AKE-M baseline 的平均横向 RMSE 降低 "
        f"{pct_reduction(f(full, 'rmse_lat_mean_mean'), f(base, 'rmse_lat_mean_mean'))}，平均纵向 RMSE 降低 "
        f"{pct_reduction(f(full, 'rmse_long_mean_mean'), f(base, 'rmse_long_mean_mean'))}，连接最大利用率降低 "
        f"{pct_reduction(f(full, 'connection_max_utilization_mean'), f(base, 'connection_max_utilization_mean'))}。"
        "这一结果支持“通信感知一致性、时延补偿和团队级安全约束降低连接误差传播”的 claim。"
        "但 TF14-full 的最大横向偏差和载荷力峰值并未优于 baseline，因此主文不能写成所有指标全面占优，也不能写成货物受力更平滑。"
    )
    doc.add_picture(str(FIGS["r0_summary"]), width=Inches(5.9))
    add_caption(doc, "图X  E0 mixed fault/noise targeted 主对比统计汇总（paired n=20）。")
    doc.add_picture(str(FIGS["r0_timeseries"]), width=Inches(5.9))
    add_caption(doc, "图X  E0 mixed fault/noise 误差时间序列对比。")

    e0c = rows(E0_COMM_SUMMARY)
    e0c_by_method = {r["method"]: r for r in e0c}
    e0c_base = e0c_by_method["baseline"]
    e0c_main = e0c_by_method["tf14_main"]
    e0c_full = e0c_by_method["tf14_phase_role"]
    e0c_zoh = e0c_by_method["zoh_consensus_surrogate"]
    add_metric_table(
        doc,
        "表X  E0 high-communication-noise 主对比（paired n=20）",
        ["方法", "成功率", "横向RMSE均值/m", "纵向RMSE均值/m", "连接最大利用率", "证书ok均值", "力峰值/N"],
        [
            ["AKE-M baseline", fmt(f(e0c_base, "success_mean"), 2), fmt(f(e0c_base, "rmse_lat_mean_mean")), fmt(f(e0c_base, "rmse_long_mean_mean")), fmt(f(e0c_base, "connection_max_utilization_mean")), fmt(f(e0c_base, "certificate_ok_ratio_mean"), 5), fmt(f(e0c_base, "force_norm_peak_mean"), 1)],
            ["TF14-main", fmt(f(e0c_main, "success_mean"), 2), fmt(f(e0c_main, "rmse_lat_mean_mean")), fmt(f(e0c_main, "rmse_long_mean_mean")), fmt(f(e0c_main, "connection_max_utilization_mean")), fmt(f(e0c_main, "certificate_ok_ratio_mean"), 5), fmt(f(e0c_main, "force_norm_peak_mean"), 1)],
            ["TF14-full", fmt(f(e0c_full, "success_mean"), 2), fmt(f(e0c_full, "rmse_lat_mean_mean")), fmt(f(e0c_full, "rmse_long_mean_mean")), fmt(f(e0c_full, "connection_max_utilization_mean")), fmt(f(e0c_full, "certificate_ok_ratio_mean"), 5), fmt(f(e0c_full, "force_norm_peak_mean"), 1)],
            ["ZOH surrogate", fmt(f(e0c_zoh, "success_mean"), 2), fmt(f(e0c_zoh, "rmse_lat_mean_mean")), fmt(f(e0c_zoh, "rmse_long_mean_mean")), fmt(f(e0c_zoh, "connection_max_utilization_mean")), fmt(f(e0c_zoh, "certificate_ok_ratio_mean"), 5), fmt(f(e0c_zoh, "force_norm_peak_mean"), 1)],
        ],
    )
    doc.add_paragraph(
        "在 high-communication-noise 主对比中，TF14-full 相对 AKE-M baseline 的平均横向 RMSE 降低 "
        f"{pct_reduction(f(e0c_full, 'rmse_lat_mean_mean'), f(e0c_base, 'rmse_lat_mean_mean'))}，平均纵向 RMSE 降低 "
        f"{pct_reduction(f(e0c_full, 'rmse_long_mean_mean'), f(e0c_base, 'rmse_long_mean_mean'))}，连接最大利用率降低 "
        f"{pct_reduction(f(e0c_full, 'connection_max_utilization_mean'), f(e0c_base, 'connection_max_utilization_mean'))}，证书 ok 均值从 "
        f"{fmt(f(e0c_base, 'certificate_ok_ratio_mean'), 5)} 提升到 {fmt(f(e0c_full, 'certificate_ok_ratio_mean'), 5)}。"
        "该结果说明在纯通信退化压力下，通信感知一致性和安全约束能显著降低连接误差传播。"
        "但 TF14-full 的力峰值相对 AKE-M 上升 "
        f"{pct_delta(f(e0c_full, 'force_norm_peak_mean'), f(e0c_base, 'force_norm_peak_mean'))}，"
        "因此受力仍应作为安全代价和后续优化目标，而不是当前优势。"
    )
    doc.add_picture(str(FIGS["r0_comm_summary"]), width=Inches(5.9))
    add_caption(doc, "图X  E0 high-communication-noise 主对比统计汇总（paired n=20）。")
    doc.add_picture(str(FIGS["r0_comm_relative"]), width=Inches(5.9))
    add_caption(doc, "图X  E0 high-communication-noise 中 TF14-full 相对 AKE-M 的指标变化；负值表示指标降低。")

    e2 = rows(E2_SUMMARY)
    e2_by_method = {r["method"]: r for r in e2}
    e2_full = e2_by_method["full_tf14"]
    e2_comm = e2_by_method["no_comm_aware"]
    e2_ppc = e2_by_method["no_ppc_progress_guard"]
    e2_zoh = e2_by_method["zoh_consensus_surrogate"]
    e2_delay = e2_by_method["no_delay_compensation"]
    e2_stable = e2_by_method["no_stable_projection"]
    add_metric_table(
        doc,
        "表X  E2 mixed fault/noise 模块消融（paired n=20）",
        ["设置", "成功率", "横向RMSE/m", "纵向RMSE/m", "连接最大利用率", "证书ok均值", "力峰值/N"],
        [
            [short_method(r["method"]), fmt(f(r, "success_mean"), 2), fmt(f(r, "rmse_lat_mean_mean")), fmt(f(r, "rmse_long_mean_mean")), fmt(f(r, "connection_max_utilization_mean")), fmt(f(r, "certificate_ok_ratio_mean"), 5), fmt(f(r, "force_norm_peak_mean"), 1)]
            for r in [e2_full, e2_comm, e2_delay, e2_stable, e2_ppc, e2_zoh]
        ],
    )
    doc.add_paragraph(
        "E2 全量消融共合并 240 条 run-level 记录，覆盖 12 个模块设置，每个设置 paired n=20。"
        "删除通信感知项后，连接最大利用率相对 TF14-full 上升 "
        f"{pct_delta(f(e2_comm, 'connection_max_utilization_mean'), f(e2_full, 'connection_max_utilization_mean'))}，"
        "纵向 RMSE 上升 "
        f"{pct_delta(f(e2_comm, 'rmse_long_mean_mean'), f(e2_full, 'rmse_long_mean_mean'))}，"
        "证书 ok 均值从 "
        f"{fmt(f(e2_full, 'certificate_ok_ratio_mean'), 5)} 降至 {fmt(f(e2_comm, 'certificate_ok_ratio_mean'), 5)}，"
        "说明通信质量加权一致性项主要通过抑制连接误差传播和降低证书裕度损失发挥作用。"
        "删除 PPC/progress guard 或采用 ZOH consensus surrogate 后成功率降至 0，纵向误差大幅放大，"
        "说明进度保护项和相较于低阶 ZOH surrogate 的闭环一致性 MPC 是保障完整路径推进的关键。"
        "同时，部分模块在该单一 mixed fault/noise 场景下呈现指标混合变化，例如删除稳定投影或单独关闭时延补偿对个别均值指标影响较小，"
        "因此时延补偿的独立贡献仍需在更高延迟强度和拓扑恢复实验中进一步验证。"
        "因此消融结论应写成模块贡献具有指标分工，而不是每个模块在所有指标上单调改善。"
    )
    doc.add_picture(str(FIGS["r2_heatmap"]), width=Inches(5.9))
    add_caption(doc, "图X  E2 模块消融热力图；误差、连接利用率和成功损失的正值表示删除对应模块后相对 TF14-full 变差，力峰值仅作为代价监测，不作为失败方法的优势解释。")
    doc.add_picture(str(FIGS["r2_effects"]), width=Inches(5.9))
    add_caption(doc, "图X  E2 关键模块消融效果对比，突出通信感知项、进度保护项和 ZOH surrogate 的失效模式。")

    e2c = rows(E2_COMM_SUMMARY)
    e2c_by_method = {r["method"]: r for r in e2c}
    e2c_full = e2c_by_method["full_tf14"]
    e2c_comm = e2c_by_method["no_comm_aware"]
    e2c_delay = e2c_by_method["no_delay_compensation"]
    e2c_zoh = e2c_by_method["zoh_consensus_surrogate"]
    add_metric_table(
        doc,
        "表X  E2 high-communication-noise 关键模块消融（paired n=20）",
        ["设置", "成功率", "横向RMSE/m", "纵向RMSE/m", "连接最大利用率", "证书ok均值", "力峰值/N"],
        [
            [short_method(r["method"]), fmt(f(r, "success_mean"), 2), fmt(f(r, "rmse_lat_mean_mean")), fmt(f(r, "rmse_long_mean_mean")), fmt(f(r, "connection_max_utilization_mean")), fmt(f(r, "certificate_ok_ratio_mean"), 5), fmt(f(r, "force_norm_peak_mean"), 1)]
            for r in [e2c_full, e2c_comm, e2c_delay, e2c_zoh]
        ],
    )
    doc.add_paragraph(
        "为避免通信韧性结论只依赖 mixed fault/noise 单一组合场景，本轮额外补充 high-communication-noise 关键消融。"
        "删除通信感知项后，连接最大利用率相对 TF14-full 上升 "
        f"{pct_delta(f(e2c_comm, 'connection_max_utilization_mean'), f(e2c_full, 'connection_max_utilization_mean'))}，"
        "纵向 RMSE 上升 "
        f"{pct_delta(f(e2c_comm, 'rmse_long_mean_mean'), f(e2c_full, 'rmse_long_mean_mean'))}，"
        "证书 ok 均值从 "
        f"{fmt(f(e2c_full, 'certificate_ok_ratio_mean'), 5)} 降至 {fmt(f(e2c_comm, 'certificate_ok_ratio_mean'), 5)}。"
        "ZOH surrogate 在该通信退化场景中路径完成率为 0，说明简单低阶保持不能替代闭环一致性 MPC。"
        "单独关闭 delay compensation 后，当前四个主指标几乎不变，因此本文只能把时延补偿写成网络感知 MPC 的组成部分，"
        "其独立收益仍需更高延迟强度、拓扑断连和恢复曲线进一步验证。"
    )
    doc.add_picture(str(FIGS["r2_comm_heatmap"]), width=Inches(5.9))
    add_caption(doc, "图X  E2 high-communication-noise 关键模块消融热力图；正值表示相对 TF14-full 变差，力峰值仅作代价监测。")
    doc.add_picture(str(FIGS["r2_comm_effects"]), width=Inches(5.9))
    add_caption(doc, "图X  E2 high-communication-noise 关键模块柱状对比，通信感知项对连接安全和证书裕度影响最强。")

    e3 = rows(E3_SUMMARY)
    add_metric_table(
        doc,
        "表X  E3 稳定性证书统计（paired n=20）",
        ["场景", "方法", "路径完成率", "证书ok比例", "最小裕度"],
        [
            [short_scenario(r["scenario"]), short_method(r["method"]), fmt(f(r, "full_path_reached_mean"), 2), fmt(f(r, "certificate_ok_ratio_mean"), 3), fmt(f(r, "certificate_min_margin_mean"), 4)]
            for r in e3
        ],
    )
    doc.add_paragraph(
        "E3 结果用于验证证书记录链和 Lyapunov 型证明假设在数值实验中的一致性。"
        "该结果不能替代理论证明，也不能单独证明优于 AKE baseline；它的作用是说明在四类场景和 20 个随机种子下，"
        "证书项、裕度项和路径完成项能够稳定落盘并与 UUB/practical boundedness 证明口径相匹配。"
    )
    doc.add_picture(str(FIGS["r3_cert"]), width=Inches(5.9))
    add_caption(doc, "图X  E3 证书统计结果（paired n=20）。")

    e7 = rows(E7_SUMMARY)
    data = []
    for scenario in ["sine_comm_noise_high", "sine_mixed_fault_noise"]:
        by_method = {r["method"]: r for r in e7 if r["scenario"] == scenario}
        b = by_method["baseline"]
        p = by_method["tf14_phase_role"]
        data.append([
            scenario,
            "AKE-M",
            fmt(f(b, "full_path_reached_mean"), 2),
            fmt(f(b, "connection_max_utilization_mean")),
            fmt(f(b, "connection_violation_count_mean"), 2),
            fmt(f(b, "force_norm_peak_mean"), 1),
        ])
        data.append([
            scenario,
            "TF14-full",
            fmt(f(p, "full_path_reached_mean"), 2),
            fmt(f(p, "connection_max_utilization_mean")),
            fmt(f(p, "connection_violation_count_mean"), 2),
            fmt(f(p, "force_norm_peak_mean"), 1),
        ])
    add_metric_table(
        doc,
        "表X  E7 载荷连接安全统计（paired n=20）",
        ["场景", "方法", "路径完成率", "连接最大利用率", "连接违规数", "力峰值/N"],
        [[short_scenario(row[0]), *row[1:]] for row in data],
    )
    doc.add_paragraph(
        "E7 结果显示，TF14-full 在通信退化和混合故障场景中均保持路径完成，连接违规数为 0，连接最大利用率相对 baseline 显著降低。"
        "这支持“连接约束满足和通信退化下相对位姿保持”的 claim。与此同时，力峰值在 TF14 中更高，说明更强的连接保持和故障重构会带来力学代价，"
        "因此力相关文字必须写为监测、审计和后续优化目标，而不是当前优势。"
    )
    doc.add_picture(str(FIGS["r7_comm"]), width=Inches(5.9))
    add_caption(doc, "图X  E7 高通信退化场景下的载荷连接安全统计。")
    doc.add_picture(str(FIGS["r7_mixed"]), width=Inches(5.9))
    add_caption(doc, "图X  E7 混合故障/通信噪声场景下的载荷连接安全统计。")

    doc.add_paragraph(
        "综上，当前版本可以主张：TF14 在 targeted mixed fault/noise 和 high-communication-noise 两组主对比中改善平均跟踪误差和连接利用率，"
        "E2 在 mixed fault/noise 与 high-communication-noise 场景下支持通信感知项、进度保护项以及相较于 ZOH surrogate 的闭环一致性 MPC 对路径完成和连接约束具有关键贡献，"
        "并在 E3/E7 多随机种子实验中形成了证书和连接安全证据链。当前版本仍不能主张：严格优于上传原文 AKE-A、"
        "所有指标全面优于 baseline、所有子模块均在所有指标上单调有效、时延补偿单独在当前消融中已被充分证明、或货物受力更平滑。"
    )

    # Move newly added body elements before the conclusion heading.
    new_elems = list(doc.element.body)[start:]
    anchor = None
    for para in doc.paragraphs:
        if para.text.strip().startswith(before_text):
            anchor = para._p
            break
    if anchor is None:
        return
    for elem in new_elems:
        anchor.addprevious(elem)
    cleanup_round10_heading(doc)


def replace_outdated_claims(doc: Document) -> None:
    replacements = {
        "与 Adaptive Koopman Embedding baseline 相比，本文不只关注单体系统的模型自适应鲁棒性，而是进一步处理多车载荷运输中的通信失配、执行亏损和团队等效力闭合问题。": (
            "与 task-aligned AKE-M mechanism comparator 相比，本文不只关注单体系统的模型自适应鲁棒性，而是进一步处理多车载荷运输中的通信失配、执行亏损和团队等效力闭合问题；严格 AKE-A 原文级复现仍作为后续 baseline 工作保留。"
        ),
        "最新 gap-closure 图包中的 E3 统计显示，nominal_koopman_mpc 与 reconfigured_ftc_mpc 模式下 certificate ok ratio 均为 1.0，且最小 contraction margin 为正。这说明当前仿真结果与上述 Lyapunov 型实用稳定性证明一致：TF14 不是保证在所有扰动下误差严格收敛为零，而是保证在模型误差、通信退化和执行器辨识误差有界时，团队误差保持在可解释、可诊断、可约束的最终有界区域内。": (
            "Round10 E3 证书统计显示，comm-high 和 mixed 场景的证书 ok 均值接近 1 且最小裕度为正，但 nominal 与 single-fault 场景仍出现证书 ok 比例不足和负裕度。因此，E3 只能作为 Lyapunov 型证书记录链与 practical boundedness 口径的数值监测证据，不能写成所有模式均严格正裕度或直接证明全局渐近稳定。"
        ),
        "当前三个模态/场景的 certificate ok ratio 为 1.0，最小裕度保持为正。": (
            "comm-high 与 mixed 场景证书裕度为正；nominal 与 single-fault 场景仍存在负裕度，证书图仅作监测证据。"
        ),
        "通信一致性与时延补偿抑制车间命令离散，故障重分配则把单车执行器亏损吸收到团队冗余中。": (
            "通信质量感知一致性主要抑制连接误差传播；时延补偿作为网络感知 MPC 的组成部分，其独立收益仍需更高延迟强度和拓扑恢复实验验证；故障重分配则把单车执行器亏损吸收到团队冗余中。"
        ),
        "因此能显著减小延迟链路造成的相位滞后。": (
            "其目标是减小延迟链路造成的相位滞后，但当前 n=20 消融尚未单独证明该项在主指标上的显著收益。"
        ),
        "相比 AKE-baseline 缺少显式链路质量映射的团队协调，TF14 把通信可靠性直接写进命令融合强度，因此在故障和高负载阶段更不容易出现车间相互拉扯。": (
            "相比 AKE-M task-aligned baseline 缺少显式链路质量映射的团队协调，TF14 把通信可靠性直接写进命令融合强度，因此在故障和高负载阶段更不容易出现连接误差传播；力峰值是否降低不能作为当前结论。"
        ),
        "车间命令离散与相位滞后被抑制，连接误差和载荷横向冲突力更小": (
            "连接误差传播被抑制；载荷受力作为代价监测项，不能写成当前统计优势"
        ),
        "结果表明，TF14 在横向和纵向跟踪精度上仍显著优于 AKE-baseline，同时把失败次数从 1 次降至 0 次。": (
            "Round10 paired n=20 targeted 主对比表明，TF14-full 在 mixed fault/noise 场景中相对 AKE-M baseline 降低平均横向 RMSE、平均纵向 RMSE 和连接最大利用率；"
            "但最大横向偏差和载荷力峰值并非优势，因此本文只主张平均跟踪和连接安全指标改善，不写所有指标全面占优。"
        ),
        "当前稿件仍需谨慎处理统计边界：E1/E2/E3/E7 已形成阶段性证据闭环，但 minimum_group_n=1，不能替代最终多 seed 显著性结论。": (
            "当前稿件仍需谨慎处理统计边界：E0 targeted、E2 全量消融、E3 和 E7 已完成 paired n=20，能够支撑主对比、模块贡献、证书统计和连接安全四类结论；"
            "但 E1 fresh Koopman 多训练种子、dose-response/topology recovery 和 AKE-A 严格复现仍未完成。"
        ),
        "最新图包显示，TF14 在高通信退化和混合故障/噪声场景下能够保持完整路径推进，并通过": (
            "Round10 多随机种子结果显示，TF14 在 targeted mixed fault/noise 主对比中改善平均跟踪误差和连接利用率，E2 支撑关键模块贡献，在 E3/E7 中形成证书与连接安全证据链；但力峰值上升表明连接保持和故障重构存在力学代价，因此结论必须限定为连接安全、模块贡献与平均误差改善。"
        ),
    }
    for para in doc.paragraphs:
        text = para.text
        for old, new in replacements.items():
            if old in text:
                para.text = text.replace(old, new)
                text = para.text
        if "当前图件仍属于单 seed gap-closure 阶段" in para.text:
            para.text = para.text.replace(
                "当前图件仍属于单 seed gap-closure 阶段，最终投稿前需补齐 n>=20 多随机种子统计。",
                "当前 Round10 已补齐 E0/E2/E3/E7 的 paired n=20 统计，可支撑 mixed/comm 主对比、模块消融、证书统计和连接安全四类有限结论；仍需补齐 E1 fresh Koopman 多训练种子、AKE-A 严格复现、强物理 DMPC baseline 和 dose-response/topology recovery。"
            )
        if "当前图包能够作为阶段性机制证据，最终稿仍需用 n>=20 多 seed 统计替换单 seed 图件。" in para.text:
            para.text = para.text.replace(
                "当前图包能够作为阶段性机制证据，最终稿仍需用 n>=20 多 seed 统计替换单 seed 图件。",
                "当前 Round10 图包中的 E0/E2/E3/E7 已具备 paired n=20 统计，可以作为主文的有限统计证据；其余单 seed 机制图仅作为补充机制说明，不能替代强 baseline、E1 fresh 和 dose/topology 证据。"
            )
        if "下一步应优先补齐 n>=20 多随机种子统计、真实 Koopman DNN 多种子重训练" in para.text:
            para.text = para.text.replace(
                "下一步应优先补齐 n>=20 多随机种子统计、真实 Koopman DNN 多种子重训练",
                "下一步应优先补齐真实 Koopman DNN 多种子重训练、AKE-A 严格复现、物理/非 Koopman DMPC baseline"
            )
        if "minimum_group_n=1" in para.text:
            para.text = (
                "本节替换旧稿中原 5.6/5.7 的 TF13 过渡实验段，统一纳入 TF14 最新图包。"
                "上一版仅能支撑代表性趋势；Round10 已补充 E0 targeted、E2 全量消融、E3 和 E7 的 paired n=20 结果，"
                "因此这些图和表可以支撑主对比、模块贡献、证书统计和连接安全四类结论。仍需注意，E1 fresh、dose-response/topology 和 AKE-A 严格复现尚未完成。"
            )
        if "Adaptive Koopman Embedding baseline" in para.text and "Singh" not in para.text:
            para.text = para.text.replace("Adaptive Koopman Embedding baseline", "task-aligned AKE-M mechanism comparator")
        if "AKE-baseline" in para.text:
            para.text = para.text.replace("AKE-baseline", "AKE-M task-aligned baseline")
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if "当前 n=1，写趋势不写显著性。" in cell.text:
                    cell.text = cell.text.replace(
                        "当前 n=1，写趋势不写显著性。",
                        "单 seed 机制图；统计结论以 Round10 paired n=20 图表为准。"
                    )
                if "车间命令离散与相位滞后被抑制，连接误差和载荷横向冲突力更小" in cell.text:
                    cell.text = cell.text.replace(
                        "车间命令离散与相位滞后被抑制，连接误差和载荷横向冲突力更小",
                        "连接误差传播被抑制；载荷受力作为代价监测项，不能写成当前统计优势"
                    )


def main() -> None:
    if not BASE.exists():
        raise FileNotFoundError(BASE)
    for path in [E0_SUMMARY, E0_COMM_SUMMARY, E2_SUMMARY, E2_COMM_SUMMARY, E3_SUMMARY, E7_SUMMARY, *FIGS.values()]:
        if not path.exists():
            raise FileNotFoundError(path)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(BASE, OUT)
    doc = Document(OUT)
    replace_outdated_claims(doc)
    insert_new_section(doc, "6 结论")
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
