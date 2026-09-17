from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DOCX = ROOT / "paper_dcn_tf12_draft" / "tf14_remaining_experiments_figures_plan_zh_2026-05-09.docx"
SKILL_SCRIPTS = Path(
    r"C:\Users\lj\.codex\plugins\cache\openai-primary-runtime\documents\26.430.10722"
    r"\skills\documents\scripts"
)
sys.path.append(str(SKILL_SCRIPTS))
from table_geometry import apply_table_geometry, column_widths_from_weights  # noqa: E402


ACCENT = RGBColor(18, 82, 112)
MUTED = RGBColor(90, 100, 110)
HEADER_FILL = "DDEFF6"
NOTE_FILL = "F5F7F9"
WARN_FILL = "FFF3CD"
GOOD_FILL = "DFF3E8"


def set_east_asia_font(run, east_asia: str = "Microsoft YaHei", latin: str = "Arial"):
    run.font.name = latin
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)


def set_style_font(style, size_pt=None, bold=None, color=None):
    style.font.name = "Arial"
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size_pt is not None:
        style.font.size = Pt(size_pt)
    if bold is not None:
        style.font.bold = bold
    if color is not None:
        style.font.color.rgb = color


def content_width_dxa(doc: Document) -> int:
    sec = doc.sections[-1]
    return int(sec.page_width.twips - sec.left_margin.twips - sec.right_margin.twips)


def shade_cell(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text: str, *, bold=False, color=None, size=7.3, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(1)
    if align is not None:
        p.alignment = align
    run = p.add_run(str(text))
    set_east_asia_font(run)
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_table(doc: Document, headers, rows, weights, font_size=7.2, header_size=7.8):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for idx, h in enumerate(headers):
        shade_cell(hdr.cells[idx], HEADER_FILL)
        set_cell_text(hdr.cells[idx], h, bold=True, color=ACCENT, size=header_size)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            set_cell_text(cells[idx], value, size=font_size)
    total = content_width_dxa(doc)
    widths = column_widths_from_weights(weights, total_width_dxa=total)
    apply_table_geometry(table, widths, table_width_dxa=total, indent_dxa=0)
    doc.add_paragraph()
    return table


def add_para(doc: Document, text: str = "", style: str | None = None, bold=False, color=None):
    p = doc.add_paragraph(style=style)
    if text:
        r = p.add_run(text)
        set_east_asia_font(r)
        r.bold = bold
        if color is not None:
            r.font.color.rgb = color
    return p


def add_bullets(doc: Document, items, size=8.5):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2.2)
        r = p.add_run(item)
        set_east_asia_font(r)
        r.font.size = Pt(size)


def add_note_box(doc: Document, title: str, body: str, fill: str = NOTE_FILL):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    shade_cell(cell, fill)
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(title)
    set_east_asia_font(r)
    r.bold = True
    r.font.size = Pt(9.2)
    r.font.color.rgb = ACCENT
    p2 = cell.add_paragraph()
    r2 = p2.add_run(body)
    set_east_asia_font(r2)
    r2.font.size = Pt(8.5)
    total = content_width_dxa(doc)
    apply_table_geometry(table, [total], table_width_dxa=total, indent_dxa=0)
    doc.add_paragraph()


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r1 = paragraph.add_run("第 ")
    set_east_asia_font(r1)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    r2 = paragraph.add_run()
    r2._r.append(fld_begin)
    r2._r.append(instr)
    r2._r.append(fld_end)
    r3 = paragraph.add_run(" 页")
    set_east_asia_font(r3)


def setup_doc() -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width = Inches(11)
    sec.page_height = Inches(8.5)
    sec.top_margin = Inches(0.55)
    sec.bottom_margin = Inches(0.5)
    sec.left_margin = Inches(0.55)
    sec.right_margin = Inches(0.55)

    styles = doc.styles
    set_style_font(styles["Normal"], 9)
    styles["Normal"].paragraph_format.line_spacing = 1.08
    styles["Normal"].paragraph_format.space_after = Pt(3)
    set_style_font(styles["Title"], 19, True, ACCENT)
    set_style_font(styles["Subtitle"], 10, False, MUTED)
    set_style_font(styles["Heading 1"], 13, True, ACCENT)
    set_style_font(styles["Heading 2"], 11, True, RGBColor(32, 73, 96))
    for name in ("List Bullet", "List Number"):
        set_style_font(styles[name], 8.6)
        styles[name].paragraph_format.left_indent = Inches(0.25)
        styles[name].paragraph_format.first_line_indent = Inches(-0.11)
        styles[name].paragraph_format.space_after = Pt(2)

    header = sec.header.paragraphs[0]
    hr = header.add_run("TF14 补充实验与出图计划 | Remaining Experiments and Figures")
    set_east_asia_font(hr)
    hr.font.size = Pt(8)
    hr.font.color.rgb = MUTED
    footer = sec.footer.paragraphs[0]
    add_page_number(footer)
    for run in footer.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = MUTED
    return doc


def build_docx() -> None:
    doc = setup_doc()

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("TF14 还需补充的实验与出图计划")
    set_east_asia_font(r)
    r.bold = True
    r.font.size = Pt(20)
    r.font.color.rgb = ACCENT
    sub = doc.add_paragraph(style="Subtitle")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = sub.add_run("面向中文稿后续补证与最终投稿统计的可执行实验清单")
    set_east_asia_font(sr)

    add_note_box(
        doc,
        "总体结论",
        "当前中文稿已经有控制效果、FDI、FTC、连接安全和实时性证据。继续补充的重点不是再堆相似轨迹图，而是补齐四个硬缺口："
        "Koopman DNN 多 seed 预测、真正逐模块消融、Lyapunov 证书有效性、phase-role 高曲率收益。另需把主结果统计从 n=6/代表性曲线升级到 n>=20。",
        WARN_FILL,
    )

    doc.add_heading("1. 优先级总览", level=1)
    add_table(
        doc,
        ["优先级", "实验编号", "必须解决的问题", "产出图", "是否主文"],
        [
            ["P0", "E0 最终多 seed 主对比", "把当前 n=6/代表曲线升级为最终统计，支撑“控制效果更好”。", "R0-1 至 R0-5", "主文"],
            ["P0", "E1 Koopman DNN 预测验证", "证明双线性 Koopman 与稳定投影本身有贡献，而不是只靠控制后处理。", "R1-1 至 R1-6", "主文+补充"],
            ["P0", "E2 真正逐模块消融", "替换 F13 available variants，证明每个模块的边际贡献。", "R2-1 至 R2-4", "主文"],
            ["P0", "E3 Lyapunov 证书修正验证", "解决 F12 中 ok=0.0%、min margin 为负的问题，避免稳定性证据反向。", "R3-1 至 R3-4", "主文/补充"],
            ["P1", "E4 通信模块消融", "证明通信质量一致性、时延补偿、约束收紧不是装饰。", "R4-1 至 R4-5", "主文+补充"],
            ["P1", "E5 FDI 分类鲁棒性", "从单例检测扩展到多故障类型分类与误报控制。", "R5-1 至 R5-5", "主文+补充"],
            ["P1", "E6 FTC severe/safe fallback", "证明安全降级分支真的能工作。", "R6-1 至 R6-5", "主文/补充"],
            ["P1", "E7 载荷连接安全公平重跑", "确认 F10 中连接改善和受力代价在同扰动条件下成立。", "R7-1 至 R7-5", "主文"],
            ["P1", "E8 phase-role 高曲率收益", "当前 DLC 收益弱，需要 hairpin/S-curve 等场景证明角色调度价值。", "R8-1 至 R8-6", "主文/补充"],
            ["P2", "E9 实时性剖析", "把 S19 从均值对比升级到 p95/p99、overrun、模块耗时。", "R9-1 至 R9-4", "补充"],
            ["P2", "E10 matched TF13 对比", "若正文要直接对比 TF13，必须在同通信/故障扰动下重跑。", "R10-1 至 R10-3", "可选主文"],
        ],
        [0.6, 1.7, 3.2, 2.0, 1.0],
        font_size=7.5,
    )

    doc.add_heading("2. 实验设计详表", level=1)
    experiments = [
        [
            "E0 最终多 seed 主对比",
            "目的：替换当前 F6 的 Stage2 n=6 统计和 S11 的代表性曲线，形成最终控制效果证据。假设：TF14 main/phase-role 在通信退化和故障场景下显著降低横向 RMSE、峰值误差和失败次数。",
            "方法：AKE/baseline、TF14 main、TF14 phase-role、TF14 error-match/all-solve；若要保留 TF13 对比，加入 matched TF13。场景：DLC clean、high communication degradation、single dual-channel fault、mixed fault/noise。",
            "规格：最终 n>=20 seeds，建议 2026-2045；所有方法共用同一随机种子、同一初值、同一故障起点/通信扰动序列。保存 time、state、ref、control、comm、fault、payload_force、connection、runtime。",
            "指标：RMSE_y/RMSE_s、max_y/max_s、full path rate、solver success、failure count、connection RMS/peak、payload force RMS/peak、step time。",
            "判定：TF14 main 或 phase-role 在 mixed/high-comm 中 RMSE_y 和 failure count 优于 baseline；若纵向误差不显著改善，需解释为横向/连接安全优先。",
        ],
        [
            "E1 Koopman DNN 预测验证",
            "目的：补 F3 的证据缺口。当前 F3 是 lightweight least-squares screening，不能证明最新 Koopman 网络优势。",
            "方法：raw linear model、linear Koopman、bilinear Koopman、bilinear+stable projection、bilinear+online adaptation；可加入历史 AKE baseline 网络作为参考。",
            "规格：离线数据固定 train/val/test；>=5 seeds，建议 10 seeds；测试轨迹覆盖 straight、DLC、hairpin、S-curve。每个模型保存 A/B/N、spectral radius、loss、one-step prediction、multi-step rollout。",
            "指标：per-state one-step RMSE、rollout error vs horizon、rollout divergence rate、spectral radius、closed-loop RMSE with model variants。",
            "判定：bilinear+stable projection 在 rollout 稳定性和控制相关状态上优于 linear/无投影；若单步误差不全优，强调长时域稳定和闭环收益。",
        ],
        [
            "E2 真正逐模块消融",
            "目的：替换 F13 available variants。审稿人会要求每个模块只关一个，其余不变。",
            "方法：full TF14、no bilinear、no stable projection、no online adaptation、no comm-aware、no delay compensation、no FDI、no FTC switching、no phase-role、no realtime scheduler、no PPC/progress guard。",
            "规格：至少 mixed fault/noise 和 high communication 两个场景；最终 n>=20 seeds。每个消融只改一个开关，不能同时改变 horizon、solver、速度、故障参数。",
            "指标：RMSE_y/RMSE_s、max_y/max_s、connection RMS、payload force peak/RMS、certificate_ok/min margin、runtime、full path rate。",
            "判定：关掉 FDI/FTC 后故障场景恶化；关掉 comm-aware 后通信场景恶化；关掉 bilinear/stable 后预测或闭环误差恶化；phase-role 至少在高曲率场景体现收益。",
        ],
        [
            "E3 Lyapunov 证书修正验证",
            "目的：解决 F12 当前 ok=0.0%、min margin 为负，不能作为正证据的问题。",
            "方法：nominal、comm high、single fault、mixed fault/noise、severe fallback 五类场景；记录每步 V_sigma、predicted upper、margin、disturbance bound、identification error、mode。",
            "规格：先检查 certificate_ok 公式是否与 practical/ISS 稳定定义一致；若理论允许非单调 V，则图中显示 bounded margin/ultimate bound，而不是强行要求每步 contraction。",
            "指标：ok ratio、bounded-margin violation count、min/mean margin、V ultimate bound、mode-wise certificate statistics。",
            "判定：证书结论必须与理论一致。若每步 contraction 不成立，就写 practical stability，并展示误差最终有界和违反次数可控。",
        ],
        [
            "E4 通信模块消融",
            "目的：突出 Digital Communications and Networks 主题，证明网络韧性模块的贡献。",
            "方法：no-comm-aware、delay-only、quality-consensus-only、constraint-tightening-only、fallback-only、full TF14。网络条件：随机 loss、burst dropout、delay jitter、asymmetric link。",
            "规格：loss rate = 0/0.1/0.25/0.4；delay = 0/1/3/5 steps；burst = 1 s/2 s/4 s；每组 >=10 seeds，重点组 n>=20。",
            "指标：q_global、pairwise q_ij、delay age、packet loss、fallback count、constraint tightening、tracking error、connection error、force peak。",
            "判定：full TF14 在高 loss/delay 下比 no-comm-aware 保持更低误差/更少约束违反；delay-only 和 quality-only 不能完全替代 full 模块。",
        ],
        [
            "E5 FDI 分类鲁棒性",
            "目的：从 F8 单例检测升级为分类和泛化证据。",
            "方法：fault types = ax degrade、steer degrade、dual degrade、severe dual loss、ramp fault、intermittent dropout、bias/stuck actuator、多车同时故障。",
            "规格：每类至少 50 个事件或 n>=20 seeds；故障车辆随机，开始 s/time 随机，故障幅值从 0.35-0.8 扫描；clean 和 high-comm 都跑。",
            "指标：detection delay、switch delay、false alarm rate、classification accuracy、precision/recall/F1、ROC/AUC、gamma estimation error。",
            "判定：FDI 在高通信扰动下仍能保持低误报，dual/severe 类故障能快速触发 FTC/safe fallback。",
        ],
        [
            "E6 FTC severe/safe fallback",
            "目的：证明 safe degraded consensus 分支真实有效，而不是只存在于方法描述。",
            "方法：full TF14、no switching、inverse compensation only、redistribution only、safe fallback disabled、oracle known-fault compensation。",
            "规格：severe dual loss、steering stuck、accelerator stuck、two-vehicle simultaneous fault、fault under high communication loss；n>=10，核心组 n>=20。",
            "指标：mode timeline、safe-mode activation time、vehicle-level command/actual input、saturation count、connection violation、payload force/moment、tracking recovery time。",
            "判定：full TF14 在 severe 场景下进入 safe/reconfigured，并减少连接违反或失败率；禁用 safe fallback 后出现更多饱和/违反/失败。",
        ],
        [
            "E7 载荷连接安全公平重跑",
            "目的：巩固 F10。当前 F10 已显示连接改善，但受力峰值有主动补偿代价，且需确认同扰动公平性。",
            "方法：same-disturbance baseline、baseline+same comm/fault、TF14 main、TF14 phase-role、no-FTC、no-comm-aware。",
            "规格：所有方法使用完全相同通信 loss/delay 序列和故障序列；DLC/high-comm/mixed/hairpin 各 n>=20。",
            "指标：Fx/Fy/Mz/norm peak、RMS、force rate、connection stretch/RMS/max utilization、violation count、corner load spread。",
            "判定：TF14 不一定降低所有力峰值，但应显著降低连接误差/违反和载荷姿态风险；正文表述为“安全约束改善”，不要写“受力全面降低”。",
        ],
        [
            "E8 phase-role 高曲率收益",
            "目的：当前 F11 在 DLC 下收益偏弱，需要用高曲率场景证明阶段-角色调度价值。",
            "方法：no phase-role、full phase-role、no inner/outer trim、no front/rear trim、weak trim、aggressive trim。",
            "规格：hairpin、S-curve、narrow passage、DLC high speed；车辆角色按 FL/FR/RL/RR 和 inner/outer、front/rear 聚合；n>=20。",
            "指标：inner/outer error gap、front/rear error gap、yaw phase area、payload yaw moment、connection RMS/peak、trim energy、control saturation。",
            "判定：phase-role 至少在高曲率场景降低角色间误差差异、连接误差或载荷 yaw moment；若平均 RMSE 改善小，也可用物理量证明机理。",
        ],
        [
            "E9 实时性剖析",
            "目的：把 S19 从均值图升级到实时控制证据。",
            "方法：decimation=1/2/5、min_solve_vehicles=1/2/4、skip_on_budget on/off、all-solve、leader-first on/off。",
            "规格：保存每步 module_time：lifting、MPC solve、FDI、FTC、phase-role、comm update、logging；n>=10。",
            "指标：step time p50/p95/p99、MPC solve p95、overrun ratio、vehicles solved、skipped solves、tracking/runtime Pareto。",
            "判定：main/phase-role 满足实时预算，all-solve/error-match 是精度上界但计算代价大；不能把 all-solve 写成实时主方法。",
        ],
        [
            "E10 matched TF13 对比",
            "目的：如果正文要直接说优于 TF13，必须补匹配场景，而不能只用 available benchmark。",
            "方法：TF13 known-fault compensation、TF14 main、TF14 phase-role、baseline/AKE。",
            "规格：同样的 high communication degradation、single fault、mixed fault/noise；故障开始、车辆、幅值一致；n>=10，最终 n>=20。",
            "指标：trajectory、RMSE、force/connection、runtime、fault response latency、full path rate。",
            "判定：若 TF13 在部分误差上更好，论文应写 TF14 的优势是未知故障自治诊断、通信韧性和安全闭环，而不是所有误差都压过 TF13。",
        ],
    ]
    add_table(
        doc,
        ["实验", "为什么补", "对比方法/场景", "运行规格", "指标", "通过标准/写作口径"],
        experiments,
        [1.25, 2.05, 2.3, 2.0, 1.8, 2.3],
        font_size=6.5,
        header_size=7.0,
    )

    doc.add_heading("3. 每组实验要出的图", level=1)
    figure_rows = [
        ["R0-1", "最终系统轨迹与局部放大", "E0", "主文", "两列代表场景：high-comm、mixed fault/noise；参考路径黑虚线，baseline 灰，TF14 main 橙，phase-role 绿；局部放大框显示偏差。", "证明不发散和轨迹保持。"],
        ["R0-2", "最终 e_y/e_s 时间序列 mean±95%CI", "E0", "主文", "2x2 面板；故障/通信窗口浅红/浅橙；横轴 time，纵轴 e_y/e_s；n>=20。", "证明控制误差稳定改善。"],
        ["R0-3", "RMSE/max/full-path 统计箱线图", "E0", "主文", "箱线+散点；每个场景一组；标 n 和 seed；显著性只在统计检验后加。", "替换当前 F6 n=6。"],
        ["R0-4", "失败次数和 solver success", "E0", "补充", "柱状图+error bar；显示 failure count、full path rate、solver success。", "证明鲁棒性不是只看 RMSE。"],
        ["R0-5", "单车误差 fault/support 聚合", "E0", "补充", "fault vehicle vs support mean；横向/纵向/航向误差分开。", "证明协同分担。"],
        ["R1-1", "全状态/输入训练测试分布", "E1", "补充", "复用 S01 但补真实 curvature 通道；train 浅橙，test 绿；标约束线。", "证明数据覆盖。"],
        ["R1-2", "DNN 训练/验证 loss mean±CI", "E1", "补充", "半对数 y 轴；每模型多 seed；显示最终 loss 分布。", "证明训练稳定。"],
        ["R1-3", "per-state one-step error", "E1", "主文/补充", "分组柱状或箱线；state = s,e_y,e_psi,vx,vy,r；方法统一配色。", "证明模型精度。"],
        ["R1-4", "multi-step rollout error", "E1", "主文", "横轴 horizon 1-30/50，纵轴 rollout error；mean±CI。", "证明长时域预测稳定性。"],
        ["R1-5", "代表 rollout 轨迹", "E1", "补充", "truth 黑，预测按方法色；DLC 和 hairpin 各一条。", "解释误差来源。"],
        ["R1-6", "spectral radius/稳定投影效果", "E1", "补充", "bar/violin；显示 projection 前后谱半径和 divergence rate。", "证明 stable projection 必要性。"],
        ["R2-1", "真实逐模块 ablation 主指标", "E2", "主文", "forest plot 或横向条形；指标为 ΔRMSE_y、Δconnection RMS、Δforce peak、Δruntime。", "证明每个模块贡献。"],
        ["R2-2", "消融时间序列代表图", "E2", "补充", "full vs no-FDI/no-FTC/no-comm/no-bilinear；故障/通信窗口标注。", "说明退化机理。"],
        ["R2-3", "消融安全指标", "E2", "主文/补充", "connection violation、certificate violation、input saturation、failure count。", "证明安全模块贡献。"],
        ["R2-4", "消融雷达/热力图", "E2", "补充", "场景 x 模块 的改善率热力图；色条 centered at 0。", "展示模块场景依赖性。"],
        ["R3-1", "修正后 certificate 时间线", "E3", "主文", "V、upper/bound、margin、mode、ok/bounded-ok；标题写 practical stability，不写单调收敛。", "修复 F12。"],
        ["R3-2", "mode-wise margin 分布", "E3", "补充", "nominal/reconfigured/safe 三组箱线图。", "证明各模式稳定边界。"],
        ["R3-3", "disturbance/identification error vs margin", "E3", "补充", "散点或二维密度；显示扰动越强 margin 越紧。", "连接理论参数。"],
        ["R3-4", "ultimate bound 验证", "E3", "主文/补充", "误差范数和理论 bound 同图；最终有界而非每步下降。", "支撑 ISS/practical stability。"],
        ["R4-1", "通信质量与跟踪误差联动", "E4", "主文", "q_global、delay/loss、e_y/e_s 同时间轴。", "证明网络扰动影响。"],
        ["R4-2", "通信模块消融柱状图", "E4", "主文", "no-comm、delay-only、quality-only、tightening-only、fallback-only、full。", "证明 full 模块必要。"],
        ["R4-3", "loss-delay 鲁棒性热力图", "E4", "补充", "x=loss，y=delay，color=RMSE/violation；每方法一张小图。", "展示鲁棒域。"],
        ["R4-4", "constraint tightening/fallback timeline", "E4", "补充", "q 低时约束收紧和 fallback 标志同步。", "证明机制触发。"],
        ["R4-5", "网络拓扑/链路质量图", "E4", "补充", "四车节点，边色表示 q_ij。", "强化 DCN 主题。"],
        ["R5-1", "FDI 多故障混淆矩阵", "E5", "主文", "蓝色单色矩阵；类别 ax/steer/dual/severe/ramp/intermittent。", "证明分类能力。"],
        ["R5-2", "FDI ROC/PR 曲线", "E5", "补充", "每类一条曲线；标 AUC。", "证明阈值鲁棒。"],
        ["R5-3", "检测/切换延迟分布", "E5", "主文/补充", "box/violin；按故障类型和通信条件分组。", "证明快速响应。"],
        ["R5-4", "gamma 估计误差", "E5", "补充", "真实效率 vs 估计效率散点，y=x 参考线。", "证明辨识准确。"],
        ["R5-5", "false alarm vs threshold 敏感性", "E5", "补充", "阈值扫描曲线。", "证明低误报不是偶然。"],
        ["R6-1", "safe fallback 模式时间线", "E6", "主文", "severe 场景；nominal/reconfigured/safe 色带。", "证明 safe 分支触发。"],
        ["R6-2", "禁用 safe 的失败对比", "E6", "主文/补充", "full vs safe-disabled 的连接违反/失败率。", "证明安全降级价值。"],
        ["R6-3", "车辆级 commanded/actual input", "E6", "补充", "u_MPC、u_FTC、u_actual 三线；饱和边界红虚线。", "证明补偿可执行。"],
        ["R6-4", "payload force under severe fault", "E6", "补充", "Fx/Fy/Mz/norm 时间序列与峰值统计。", "证明严重故障下安全。"],
        ["R6-5", "recovery time 统计", "E6", "补充", "故障后误差回到阈值内所需时间。", "证明恢复能力。"],
        ["R7-1", "同扰动连接误差对比", "E7", "主文", "baseline/no-comm-aware/TF14 同扰动；connection stretch/RMS。", "巩固 F10。"],
        ["R7-2", "force RMS/peak tradeoff", "E7", "主文", "峰值和 RMS 分开；不要只画 peak。", "解释力峰值代价。"],
        ["R7-3", "connection violation count", "E7", "主文/补充", "每场景 violation count 和 max utilization。", "证明安全约束。"],
        ["R7-4", "corner load spread", "E7", "补充", "四角载荷 spread 时间序列/统计。", "证明载荷均衡。"],
        ["R7-5", "force rate/jerk", "E7", "补充", "dF/dt 或控制变化率。", "避免主动补偿过激。"],
        ["R8-1", "高曲率路径阶段标签", "E8", "主文/补充", "hairpin/S-curve 上标 straight/entry/core/exit。", "证明 phase 定义。"],
        ["R8-2", "phase-role trim 高曲率图", "E8", "主文", "delta_trim/ax_trim 四车线，阶段背景色。", "证明调度器工作。"],
        ["R8-3", "inner/outer error gap", "E8", "主文", "按内外侧聚合，full phase-role vs no phase-role。", "证明角色收益。"],
        ["R8-4", "front/rear load/error split", "E8", "补充", "前后角色聚合误差和载荷。", "证明前后角色。"],
        ["R8-5", "yaw phase area", "E8", "补充", "e_psi-r phase plane 面积统计。", "证明动态稳定性。"],
        ["R8-6", "payload yaw moment reduction", "E8", "主文/补充", "Mz RMS/peak 对比。", "证明载荷机理。"],
        ["R9-1", "runtime Pareto final", "E9", "补充/主文可选", "x=step p95，y=RMSE_y/connection RMS，点大小=overrun。", "证明实时折中。"],
        ["R9-2", "step time 分布", "E9", "补充", "hist/violin；预算线 20 ms 或配置预算。", "证明实时风险。"],
        ["R9-3", "模块耗时堆叠图", "E9", "补充", "lifting/MPC/FDI/FTC/phase/comm/logging 堆叠。", "定位计算瓶颈。"],
        ["R9-4", "vehicles solved/skipped", "E9", "补充", "每步求解车辆数和 skip 标志。", "证明 scheduler 行为。"],
        ["R10-1", "matched TF13 轨迹+误差", "E10", "主文可选", "同场景同扰动；TF13 蓝，TF14 橙/绿。", "公平对比 TF13。"],
        ["R10-2", "matched TF13 故障响应", "E10", "补充", "known-fault TF13 vs online FDI TF14 的检测/切换/补偿链。", "突出未知故障自治。"],
        ["R10-3", "matched TF13 物理安全指标", "E10", "主文可选", "force/connection/certificate/runtime 一组图。", "证明优势不是只靠 RMSE。"],
    ]
    add_table(
        doc,
        ["图号", "图名", "实验", "位置", "画法规格", "证明点"],
        figure_rows,
        [0.65, 2.0, 0.65, 0.85, 4.0, 2.0],
        font_size=6.6,
        header_size=7.1,
    )

    doc.add_heading("4. 改进点到实验的对应关系", level=1)
    add_table(
        doc,
        ["改进点", "必须验证的问题", "对应实验", "核心图", "当前状态"],
        [
            ["双线性 Koopman 与稳定投影", "预测更准/更稳，闭环不是靠后处理硬撑。", "E1、E2", "R1-3、R1-4、R1-6、R2-1", "未闭环，P0。"],
            ["通信质量一致性、时延补偿、约束收紧", "通信差时是否真的比无通信增强好。", "E4、E0", "R4-1、R4-2、R4-3、R0-2", "部分覆盖，P1。"],
            ["在线 FDI 与效率辨识", "未知故障能否快速、低误报、分类准确。", "E5", "R5-1、R5-3、R5-4", "单例覆盖，需分类。"],
            ["FTC 切换与重分配", "故障后切换是否有效，重分配是否可执行。", "E6、E2", "R6-1、R6-2、R6-3、R2-1", "普通故障覆盖，severe 缺。"],
            ["载荷连接安全", "连接误差/违反是否降低，力峰值代价是否可接受。", "E7、E0", "R7-1、R7-2、R7-3", "需公平多场景。"],
            ["phase-role 调度", "高曲率/角色差异下是否有收益。", "E8", "R8-2、R8-3、R8-6", "DLC 收益弱，需高曲率。"],
            ["Lyapunov/practical stability certificate", "证书定义和实验统计是否一致。", "E3", "R3-1、R3-2、R3-4", "F12 当前不可作为正证据。"],
            ["实时预算", "主方法是否满足实时，all-solve 的代价是什么。", "E9", "R9-1、R9-2、R9-3", "均值覆盖，p95/p99 缺。"],
            ["完整消融", "每个模块的边际贡献。", "E2", "R2-1、R2-3", "F13 不是真消融，P0。"],
        ],
        [2.0, 2.4, 1.1, 2.0, 2.2],
        font_size=7.2,
    )

    doc.add_heading("5. 推荐执行顺序", level=1)
    add_bullets(
        doc,
        [
            "第一批必须先跑：E0、E1、E2、E3。它们分别对应最终主结果、Koopman 贡献、完整消融和稳定性证书，是审稿最容易卡住的四个点。",
            "第二批随后跑：E4、E5、E6、E7。它们会把 DCN 主题、FDI 分类、FTC safe fallback 和载荷安全做实。",
            "第三批增强说服力：E8、E9、E10。phase-role 高曲率、实时性剖析和 matched TF13 对比是加分项，但不要抢在 P0 缺口前面。",
            "所有 P0/P1 实验统一保存 manifest.json、config.json、source CSV/NPZ、plot 脚本、PDF/SVG/PNG，避免后续图能看但无法复现。",
            "最终稿写作时，结论强度要按证据等级来：n>=20 多 seed 统计可以写强结论；代表性曲线只能写案例；available variants 不能写成 ablation。",
        ],
    )

    add_note_box(
        doc,
        "最小补证包",
        "如果时间非常紧，最少也要完成：E0 最终 n>=20 主对比、E1 Koopman DNN 多 seed 预测、E2 真正逐模块消融、E3 稳定性证书修正、E7 同扰动载荷连接安全。"
        "这五组完成后，论文的“控制效果好、模块有贡献、理论证书不反向、协同运输安全性改善”四条主线才算站住。",
        GOOD_FILL,
    )

    doc.save(OUT_DOCX)


if __name__ == "__main__":
    build_docx()
    print(OUT_DOCX)
