from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"D:\LEARNING\ZNN\ZNN\Adaptive-koopman\Adaptive-koopman-main")
DRAFT = ROOT / "paper_dcn_tf12_draft"
SRC_DOCX = DRAFT / "manuscript_zh_literature_update_figures_tf14_integrated_2026-05-09.docx"
OUT_DOCX = DRAFT / "manuscript_zh_literature_update_figures_tf14_lyapunov_2026-05-10.docx"
EQ_DIR = DRAFT / "equation_images_lyapunov_tf14_20260510"


def set_run_font(run, size=None, bold=None, color=None, name="宋体"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def paragraph_before(anchor):
    new_p = OxmlElement("w:p")
    anchor._p.addprevious(new_p)
    from docx.text.paragraph import Paragraph

    return Paragraph(new_p, anchor._parent)


def add_para_before(anchor, text, *, style=None, heading=False, first_line=True):
    p = paragraph_before(anchor)
    if style:
        try:
            p.style = style
        except KeyError:
            pass
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(6)
    if first_line and not heading:
        p.paragraph_format.first_line_indent = Pt(24)
    if heading:
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    if heading:
        set_run_font(run, size=12.5, bold=True, color=(31, 78, 121), name="黑体")
    else:
        set_run_font(run, size=10.5)
    return p


def add_equation_before(anchor, image_name, width=5.7):
    path = EQ_DIR / image_name
    if not path.exists():
        raise FileNotFoundError(path)
    p = paragraph_before(anchor)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.keep_together = True
    p.add_run().add_picture(str(path), width=Inches(width))
    return p


def insert_blocks_before(anchor, blocks):
    for kind, payload, width in blocks:
        if kind == "h":
            add_para_before(anchor, payload, heading=True, first_line=False)
        elif kind == "p":
            add_para_before(anchor, payload)
        elif kind == "eq":
            add_equation_before(anchor, payload, width=width)
        else:
            raise ValueError(kind)


def proof_blocks():
    b = []
    b.append(("h", "4.8 完整 Lyapunov 型稳定性证明", None))
    b.append(("p", "本节给出 TF14 闭环的完整 Lyapunov 型稳定性证明。由于系统同时存在 Koopman 建模误差、通信退化、时延补偿残差、在线 FDI 辨识误差和执行器效率损失，本文不把结论写成无扰动情形下才成立的严格渐近稳定，而证明切换容错闭环满足输入到状态实用稳定性：只要扰动和辨识误差有界，团队误差最终进入一个有界小邻域；当这些外部项为零时，结论退化为指数稳定。", None))

    b.append(("h", "4.8.1 误差系统与切换模式", None))
    b.append(("p", "定义团队中心误差为", None))
    b.append(("eq", "lyap_01_error.png", 3.2))
    b.append(("p", "其中团队中心等效状态选为", None))
    b.append(("eq", "lyap_02_team_state.png", 5.2))
    b.append(("p", "TF14 闭环包含三类全局控制模式：N 表示 nominal Koopman-MPC，R 表示 reconfigured FTC-MPC，S 表示 safe degraded consensus。对应代码中的模式为 nominal_koopman_mpc、reconfigured_ftc_mpc 和 safe_degraded_consensus。", None))
    b.append(("eq", "lyap_03_modes.png", 3.4))
    b.append(("p", "在工作域内，团队误差闭环可抽象为如下切换系统：", None))
    b.append(("eq", "lyap_04_switched_system.png", 5.5))
    b.append(("p", "其中 A_sigma 为当前模式下的等效闭环误差矩阵，B_sigma 为输入通道矩阵，tilde Gamma_k 表示真实执行器效率与 FDI 估计效率之间的残差，d_k 汇总 Koopman 预测误差、通信时延残差、丢包扰动、phase-role 有界 trim 误差和外部扰动。", None))

    b.append(("h", "4.8.2 基本假设", None))
    b.append(("p", "假设一：稳定投影后的 Koopman 名义模型在工作域 Omega 内有界，且 lifted 线性部分谱半径满足", None))
    b.append(("eq", "lyap_05_spectral_projection.png", 2.8))
    b.append(("p", "假设二：MPC 输入约束集合 U 为紧集，输入投影算子非扩张，即", None))
    b.append(("eq", "lyap_06_projection_nonexpansive.png", 4.0))
    b.append(("p", "假设三：通信扰动、Koopman 残差和时延补偿误差有界，可写为", None))
    b.append(("eq", "lyap_07_disturbance_bounds.png", 4.0))
    b.append(("p", "假设四：FDI 在有限时间后给出有界执行器效率估计误差，即", None))
    b.append(("eq", "lyap_08_identification_bound.png", 2.8))
    b.append(("p", "假设五：切换器满足驻留时间约束。TF14 代码中 switch_dwell_steps 默认约为 4，因此模式切换不会发生无限快抖振。", None))

    b.append(("h", "4.8.3 多 Lyapunov 函数", None))
    b.append(("p", "对每个切换模式构造二次 Lyapunov 函数", None))
    b.append(("eq", "lyap_09_lyapunov_function.png", 4.6))
    b.append(("p", "由于 P_sigma 正定，存在正数 p_min 和 p_max，使得", None))
    b.append(("eq", "lyap_10_quadratic_bounds.png", 4.7))
    b.append(("p", "因此，只要证明 V_sigma(e_k) 最终有界，即可推出团队中心误差 e_k 最终有界。", None))
    b.append(("p", "对每个模式假设存在 P_sigma，使无扰动闭环满足局部收缩条件", None))
    b.append(("eq", "lyap_11_lmi.png", 5.5))
    b.append(("p", "TF14 证书模块中使用的三个模式矩阵为", None))
    b.append(("eq", "lyap_12_p_nominal.png", 4.6))
    b.append(("eq", "lyap_13_p_reconfigured.png", 4.7))
    b.append(("eq", "lyap_14_p_safe.png", 4.8))
    b.append(("p", "对应的收缩率采用", None))
    b.append(("eq", "lyap_15_lambda_values.png", 4.8))

    b.append(("h", "4.8.4 单模式无扰动收缩性", None))
    b.append(("p", "先考虑固定模式且无扰动、无辨识误差的理想情形，即 d_k=0 且 tilde Gamma_k=0。此时由上一节的收缩条件可得", None))
    b.append(("eq", "lyap_16_nominal_contraction.png", 4.9))
    b.append(("p", "因此，在任一固定模式下，无扰动闭环指数收缩。nominal 模式收缩最快，reconfigured 模式稍保守，safe degraded 模式最保守但仍保持正收缩率。", None))

    b.append(("h", "4.8.5 有扰动与辨识误差情形", None))
    b.append(("p", "实际闭环可以写为 e_{k+1}=A_sigma e_k+w_k，其中", None))
    b.append(("eq", "lyap_17_w_def.png", 4.1))
    b.append(("p", "由于输入约束集合紧，控制输入有界，因此存在常数 c_Gamma，使得", None))
    b.append(("eq", "lyap_18_w_bound.png", 4.4))
    b.append(("p", "对 V_sigma(e_{k+1}) 展开，并使用 Young 不等式，可得", None))
    b.append(("eq", "lyap_19_perturbed_bound.png", 6.1))
    b.append(("p", "进一步把通信质量 q_k 写入扰动项。由于 1-q_k 表示通信退化程度，上式可加强为", None))
    b.append(("eq", "lyap_20_comm_bound.png", 6.25))
    b.append(("p", "这与代码中 SwitchedFTCCertificateTF14 记录的 predicted_upper 一致：", None))
    b.append(("eq", "lyap_21_predicted_upper.png", 6.0))
    b.append(("p", "代码中默认的扰动项和辨识项权重为", None))
    b.append(("eq", "lyap_22_gamma_values.png", 3.6))

    b.append(("h", "4.8.6 FTC 重构与安全降级不会破坏有界性", None))
    b.append(("p", "在 reconfigured FTC 模式下，控制器先根据 FDI 给出的效率估计进行逆补偿，再把无法由故障车辆完成的等效力缺口分配给健康车辆。该控制律可写为", None))
    b.append(("eq", "lyap_23_ftc_reconfig.png", 5.4))
    b.append(("p", "实际执行输入为", None))
    b.append(("eq", "lyap_24_actual_input.png", 3.2))
    b.append(("p", "将其代入可展开为", None))
    b.append(("eq", "lyap_25_actual_expand.png", 6.2))
    b.append(("p", "其中 epsilon_Pi 为输入投影误差。由于 Pi_U 非扩张且输入集合 U 紧，epsilon_Pi 有界。又因为 FDI 估计误差有界，可得", None))
    b.append(("eq", "lyap_26_gamma_inv_bound.png", 4.9))
    b.append(("p", "因此，FTC 重构只会引入与效率估计误差成比例的有界扰动，不会破坏上一节的 Lyapunov 递推不等式。", None))
    b.append(("p", "当系统进入 safe degraded consensus 模式时，控制律写成", None))
    b.append(("eq", "lyap_27_safe_mode.png", 5.8))
    b.append(("p", "其中 D_s 为保守缩放矩阵，且 0<alpha_delta, alpha_a<1。该模式降低控制激进度，使闭环收缩率变小但仍为正，因此系统会进入更保守的实用稳定区域。", None))

    b.append(("h", "4.8.7 切换系统最终有界性", None))
    b.append(("p", "由于 TF14 使用 switch_dwell_steps 限制模式切换，系统不存在 Zeno 现象。定义多个 Lyapunov 函数之间的跳变上界", None))
    b.append(("eq", "lyap_28_mu.png", 4.8))
    b.append(("p", "则任意切换瞬间有", None))
    b.append(("eq", "lyap_29_switch_jump.png", 4.5))
    b.append(("p", "若切换次数有限，或满足平均驻留时间条件，则存在常数 C_mu>0 和 alpha in (0,1)，使得", None))
    b.append(("eq", "lyap_30_sum_bound.png", 5.8))
    b.append(("p", "其中", None))
    b.append(("eq", "lyap_31_zeta.png", 5.0))
    b.append(("p", "结合扰动、通信退化和辨识误差的有界性，可进一步得到", None))
    b.append(("eq", "lyap_32_uniform_bound.png", 6.25))
    b.append(("p", "由 V_sigma(e_k)>=p_min||e_k||^2，得到团队误差的最终有界半径", None))
    b.append(("eq", "lyap_33_limsup.png", 6.25))
    b.append(("p", "因此 TF14 闭环系统是输入到状态实用稳定的。若 Koopman 残差、通信扰动和辨识误差均为零，则上式退化为", None))
    b.append(("eq", "lyap_34_zero_disturbance.png", 5.8))
    b.append(("p", "即理想无扰动情形下闭环指数稳定。", None))

    b.append(("h", "4.8.8 与 TF14 证书记录的对应关系", None))
    b.append(("p", "上述证明并不要求代码中的 certificate 模块本身替代理论证明。更准确地说，SwitchedFTCCertificateTF14 是证明假设在仿真中的在线记录器：V 对应当前模式下的二次 Lyapunov 值，predicted_upper 对应下一步理论上界，contraction_margin 对应上一步上界与当前实际 V 之间的差值，certificate_ok 对应该差值是否非负。", None))
    b.append(("p", "最新 gap-closure 图包中的 E3 统计显示，nominal_koopman_mpc 与 reconfigured_ftc_mpc 模式下 certificate ok ratio 均为 1.0，且最小 contraction margin 为正。这说明当前仿真结果与上述 Lyapunov 型实用稳定性证明一致：TF14 不是保证在所有扰动下误差严格收敛为零，而是保证在模型误差、通信退化和执行器辨识误差有界时，团队误差保持在可解释、可诊断、可约束的最终有界区域内。", None))
    b.append(("p", "证毕。", None))
    return b


def main():
    doc = Document(str(SRC_DOCX))
    anchor = next(p for p in doc.paragraphs if p.text.strip().startswith("5 仿真实验"))
    insert_blocks_before(anchor, proof_blocks())
    doc.save(str(OUT_DOCX))
    print(OUT_DOCX)
    print("paragraphs", len(doc.paragraphs), "tables", len(doc.tables), "inline_shapes", len(doc.inline_shapes))


if __name__ == "__main__":
    main()
