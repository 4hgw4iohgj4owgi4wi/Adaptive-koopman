from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"D:\LEARNING\ZNN\ZNN\Adaptive-koopman\Adaptive-koopman-main")
DRAFT = ROOT / "paper_dcn_tf12_draft"
BASE_DOCX = DRAFT / "manuscript_zh_literature_update_figures_2026-05-08.docx"
OUT_DOCX = DRAFT / "manuscript_zh_literature_update_figures_tf14_integrated_2026-05-09.docx"
EQ_DIR = DRAFT / "equation_images_20260509"
FIG_MAIN = ROOT / "tf14_paper_figures_20260508" / "figures"
FIG_REMAIN = ROOT / "tf14_remaining_experiments_20260509" / "figures"
FIG_GAP = ROOT / "tf14_final_gap_closure_20260509" / "figures"


def set_run_font(run, size=None, bold=None, color=None, name="宋体"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def set_paragraph_text(paragraph, text, size=10.5):
    paragraph.text = ""
    run = paragraph.add_run(text)
    set_run_font(run, size=size)
    paragraph.paragraph_format.line_spacing = 1.15
    paragraph.paragraph_format.space_after = Pt(6)


def set_title(paragraph, text, size=16):
    paragraph.text = ""
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=True, color=(31, 78, 121), name="黑体")


def replace_text_everywhere(doc, replacements):
    def replace_in_para(p):
        text = p.text
        new = text
        for old, value in replacements.items():
            new = new.replace(old, value)
        if new != text:
            style_name = p.style.name
            p.text = ""
            run = p.add_run(new)
            if style_name.startswith("Heading"):
                set_run_font(run, name="黑体", bold=True)
            elif style_name == "Title":
                set_run_font(run, size=16, bold=True, color=(31, 78, 121), name="黑体")
            else:
                set_run_font(run, size=10.5)

    for p in doc.paragraphs:
        replace_in_para(p)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    replace_in_para(p)


def add_paragraph_before(anchor, text, style=None, size=10.5, first_line=True):
    new_p = OxmlElement("w:p")
    anchor._p.addprevious(new_p)
    from docx.text.paragraph import Paragraph

    paragraph = Paragraph(new_p, anchor._parent)
    if style:
        paragraph.style = style
    if first_line and not (style and style.startswith("Heading")):
        paragraph.paragraph_format.first_line_indent = Pt(24)
    paragraph.paragraph_format.line_spacing = 1.15
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(text)
    if style and style.startswith("Heading"):
        set_run_font(run, bold=True, color=(31, 78, 121), name="黑体")
    else:
        set_run_font(run, size=size)
    return paragraph


def move_before_anchor(anchor, element):
    anchor._p.addprevious(element)


def add_body(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(24)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    set_run_font(run, size=10.5)
    return p


def add_heading(doc, text, level=2):
    p = doc.add_paragraph()
    style_candidates = [f"Heading {level}", f"标题 {level}"]
    for candidate in style_candidates:
        try:
            p.style = candidate
            break
        except KeyError:
            continue
    for run in p.runs:
        set_run_font(run, bold=True, color=(31, 78, 121), name="黑体")
    run = p.add_run(text)
    set_run_font(run, bold=True, color=(31, 78, 121), name="黑体")
    return p


def add_equation(doc, image_name, width=5.65):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(7)
    p.add_run().add_picture(str(EQ_DIR / image_name), width=Inches(width))
    return p


def add_figure(doc, path, caption, width=6.15):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Inches(width))

    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.space_after = Pt(8)
    c.paragraph_format.keep_together = True
    run = c.add_run(caption)
    set_run_font(run, size=9, color=(80, 80, 80))
    return [p, c]


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_twips):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_twips))
    tc_w.set(qn("w:type"), "dxa")


def set_table_width(table, widths):
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        table._tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)


def coerce_docx_width_numbers(doc):
    """Some source tables carry Word widths as decimal strings such as 0.0."""
    for element in doc.element.iter():
        for attr_name in (qn("w:w"),):
            value = element.get(attr_name)
            if value is None:
                continue
            try:
                if "." in value:
                    element.set(attr_name, str(int(float(value))))
            except ValueError:
                continue


def normalize_source_tables(doc):
    source_widths = [
        [2300, 7060],
        [2100, 2200, 5060],
        [1450, 1900, 1900, 2000, 2110],
        [1850, 2500, 2500, 2510],
        [1350, 1500, 1500, 1500, 1500, 2010],
        [2600, 2250, 2250, 2260],
        [2600, 3380, 3380],
    ]
    for table, widths in zip(doc.tables[:8], source_widths):
        set_table_width(table, widths)


def normalize_sections(doc):
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)


def set_cell_text(cell, text, bold=False, fill=None, color=(34, 34, 34), center=False):
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    if fill:
        shade_cell(cell, fill)
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.06
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run_font(run, size=8.6, bold=bold, color=color)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    for candidate in ["Table Grid", "网格型"]:
        try:
            table.style = candidate
            break
        except KeyError:
            continue
    for cell, header in zip(table.rows[0].cells, headers):
        set_cell_text(cell, header, bold=True, fill="EAF2F8", color=(31, 78, 121), center=True)
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            set_cell_text(cell, value)
    set_table_width(table, widths)
    doc.add_paragraph()
    return table


def build_method_update_blocks(doc):
    blocks = []
    blocks.append(add_heading(doc, "4.7 TF14 新增：在线 FDI、切换 FTC、阶段角色调度与证书记录", level=2))
    blocks.append(add_body(doc, "相对于前一版只在已知故障注入后做补偿的写法，TF14 的核心变化是把未知执行器退化的在线发现、效率估计、切换容错和安全证书记录纳入闭环。该部分不替代前述 Koopman-MPC，而是放在 Koopman 预测和通信一致性之后，作为从“名义控制命令”到“实际可执行团队命令”的韧性执行层。"))
    blocks.append(add_body(doc, "在线 FDI 模块对每辆车执行一阶预测，并比较预测状态与实测状态，得到归一化残差与 EWMA 诊断量："))
    blocks.append(add_equation(doc, "eq_fdi_residual.png"))
    blocks.append(add_equation(doc, "eq_fdi_ewma.png"))
    blocks.append(add_body(doc, "当前阈值设置包括 residual_threshold=0.85、CUSUM 阈值 8.0、纵向效率故障阈值 0.78、转向效率故障阈值 0.80、严重退化阈值 0.35。检测需要 hold_steps=5，释放需要 release_hold_steps=10，以避免把通信抖动或短时模型误差误判为物理执行器故障。执行器效率被写成"))
    blocks.append(add_equation(doc, "eq_actuator_efficiency.png", width=5.1))
    blocks.append(add_body(doc, "其中 Gamma_i,k 描述转向和纵向加速度两个通道的实际执行能力。FDI 输出不是简单报警，而是直接决定 FTC 控制模式：nominal_koopman_mpc、reconfigured_ftc_mpc 和 safe_degraded_consensus。故障置信度超过阈值时，控制器先对故障车进行效率逆补偿，并把剩余团队等效力缺口分配给健康车辆："))
    blocks.append(add_equation(doc, "eq_ftc_reconfig.png", width=5.95))
    blocks.append(add_body(doc, "当双执行器严重退化或通信质量过低时，控制器不再追求激进补偿，而是进入安全降级一致性控制："))
    blocks.append(add_equation(doc, "eq_ftc_safe.png", width=5.55))
    blocks.append(add_body(doc, "PhaseRoleSchedulerTF14 则在 FTC 输出后加入有界小修正。它根据 straight、turn_entry、turn_core、turn_exit 和 turn_transition 等路径阶段，以及 FL、FR、RL、RR 的内外侧和前后角色，调节每辆车的转角与加速度微修正："))
    blocks.append(add_equation(doc, "eq_phase_role_output.png", width=5.55))
    blocks.append(add_body(doc, "稳定性证明采用切换系统的输入到状态实用稳定性表述。令 sigma_k 属于 nominal、reconfigured 和 safe degraded 三类模式，则闭环误差可抽象为"))
    blocks.append(add_equation(doc, "eq_switched_error.png", width=4.8))
    blocks.append(add_body(doc, "在平均驻留时间、扰动有界和效率估计误差有界条件下，构造 V_sigma(e)=e^T P_sigma e，可得到如下 Lyapunov 型不等式："))
    blocks.append(add_equation(doc, "eq_lyapunov_iss.png", width=5.95))
    blocks.append(add_body(doc, "代码中的 SwitchedFTCCertificateTF14 在每一步记录 V、predicted_upper、contraction_margin 和 certificate_ok。它的作用是把理论假设和仿真实验连接起来：若证书裕度长期为正且 certificate_ok 保持为真，说明切换 FTC 并未破坏闭环实用稳定性；若裕度下降，则应触发更保守的安全降级或约束收紧。"))
    return blocks


def build_experiment_update_blocks(doc):
    blocks = []
    blocks.append(add_heading(doc, "5.6 TF14 最新图包与阶段性证据链", level=2))
    blocks.append(add_body(doc, "本节替换旧稿中原 5.6/5.7 的 TF13 过渡实验段，统一纳入 TF14 最新图包。需要特别说明：这些图已经能支撑“控制效果、在线诊断、容错切换、稳定性证书、力学安全诊断”的阶段性证据链，但 final gap-closure 包仍标注 minimum_group_n=1。因此本文当前版本可以写趋势、机制和代表性闭环结果，最终投稿前仍需补 n>=20 多随机种子统计。"))

    rows = [
        ("在线 FDI", "fig08_fdi_identification", "支撑未知执行器退化可被检测、分类和置信度跟踪。", "可作为正文机制图。"),
        ("FTC 切换", "fig09_ftc_switch_redistribution", "支撑 nominal 到 reconfigured FTC 的模式转移与补偿重分配。", "可作为正文机制图。"),
        ("phase-role", "fig11_phase_role_scheduler", "说明不同路径阶段和车辆角色的有界 trim。", "建议正文或补充材料。"),
        ("Lyapunov 证书", "fig12 / E3", "把 V、证书裕度、practical bound 和模式切换连到实验。", "正文放一张，补充材料放统计。"),
        ("消融热力图", "E2 positive=worse", "正值代表去掉模块后相对 TF14 主方法变差，读者不需要反向解释。", "当前 n=1，写趋势不写显著性。"),
        ("力学安全", "E7 force-rate/jerk/corner-load", "说明 TF14 的恢复能力伴随力峰值和 jerk 代价，需要写成权衡。", "适合补充材料或正文讨论。"),
    ]
    blocks.append(add_table(doc, ["模块", "图件", "能证明什么", "入文判断"], rows, [1500, 2200, 3800, 1860]))

    figures = [
        (FIG_REMAIN / "R0_error_timeseries.png", "图 18  TF14 在高通信退化与混合故障/噪声下的横向和纵向误差时间序列。该图用于补强系统级控制效果，但仍属于单 seed 代表性结果。"),
        (FIG_MAIN / "fig08_fdi_identification.png", "图 19  TF14 在线 FDI 与执行器效率辨识。残差、效率估计、置信度和故障类别形成完整诊断链。"),
        (FIG_MAIN / "fig09_ftc_switch_redistribution.png", "图 20  TF14 FTC 切换与补偿重分配。该图说明故障后控制器从名义 MPC 转入 reconfigured FTC，并对健康车辆进行支持补偿。"),
        (FIG_MAIN / "fig11_phase_role_scheduler.png", "图 21  阶段-角色调度器输出。不同路径阶段下的角色化 trim 说明该模块并非普通规则控制，而是显式编码车辆在载荷运输中的前后/内外侧作用。"),
        (FIG_GAP / "E1_koopman_prediction_validation.png", "图 22  Koopman 预测验证与稳定投影。该图说明双线性项的一阶拟合收益、长时开环外推风险以及稳定投影的谱半径约束效果，不能替代多 seed DNN 重训练。"),
        (FIG_GAP / "E2_positive_worse_heatmap_dlc_comm_noise_high.png", "图 23  高通信退化场景下的 positive=worse 消融热力图。红色表示去掉模块后相对 TF14 主方法变差。"),
        (FIG_GAP / "E2_positive_worse_heatmap_dlc_mixed_fault_noise.png", "图 24  混合故障/噪声场景下的 positive=worse 消融热力图。no_comm_aware、no_PPC_progress_guard 以及 FDI/FTC 相关项是主要退化来源。"),
        (FIG_GAP / "E3_certificate_margin_bound_distribution.png", "图 25  模态证书裕度与 practical-bound 统计。当前三个模态/场景的 certificate ok ratio 为 1.0，最小裕度保持为正。"),
        (FIG_GAP / "E7_force_rate_jerk_corner_timeline_dlc_mixed_fault_noise.png", "图 26  混合故障/噪声场景下的合力变化率、jerk 与角点载荷时间线。该图应写成故障恢复能力与力学代价之间的权衡，而不是“所有力学指标都优于 baseline”。"),
    ]
    for path, caption in figures:
        blocks.extend(add_figure(doc, path, caption))

    rows2 = [
        ("E1 Koopman", "一阶预测：bilinear 略优；60-step 开环 rollout 暴露外推风险；谱半径由约 1.006 压到 0.999。", "可以写稳定投影和闭环 MPC 保护必要，不能写长时开环预测全面胜出。"),
        ("E2 消融", "高通信退化下 no_comm_aware 的 conn RMS 退化最明显；混合故障下 no_PPC_progress_guard 与 no_comm_aware 退化突出。", "可以写模块去除后的趋势性退化，不能写统计显著。"),
        ("E3 证书", "三个模态/场景 ok ratio 均为 1.0，最小裕度为正。", "可以作为 Lyapunov/ISS 证明的仿真闭环证据。"),
        ("E7 力学安全", "相对 baseline，TF14 故障恢复可能带来更高 force-rate/jerk；相对 no_ftc_switching，混合故障下 TF14 main 降低力峰值、rate、jerk 与角点载荷峰值。", "应写成可解释、可约束的恢复过程，而不是无代价改善。"),
    ]
    blocks.append(add_table(doc, ["证据项", "当前定量读法", "论文写法边界"], rows2, [1600, 4850, 2910]))
    blocks.append(add_body(doc, "因此，旧稿中原 5.6 和 5.7 的复杂路径、通信消融与载荷安全叙事已由本节 TF14 图包替换。正式投稿时建议继续用 TF14 多 seed 统计图更新本节，并保留这里的审稿口径：机制可以先讲清楚，统计显著性必须等最终实验补齐。"))
    return blocks


def move_blocks_before(anchor, blocks):
    for block in blocks:
        element = block._p if hasattr(block, "_p") else block._tbl
        move_before_anchor(anchor, element)


def remove_blocks_between(start_para, end_para):
    parent = start_para._p.getparent()
    current = start_para._p
    end = end_para._p
    while current is not None and current is not end:
        nxt = current.getnext()
        parent.remove(current)
        current = nxt


def main():
    doc = Document(str(BASE_DOCX))

    set_title(doc.paragraphs[0], "基于网络韧性双线性 Koopman 学习与切换容错协同 MPC 的四车载荷运输控制", size=16)
    set_title(doc.paragraphs[1], "Network-Resilient Cooperative Transport Control via Bilinear Koopman Learning and Delay-Compensated Consensus MPC", size=12)
    set_paragraph_text(
        doc.paragraphs[3],
        "本文面向存在通信时延、丢包、未知执行器退化和刚性载荷耦合的四车协同运输任务，提出一种 TF14 网络韧性协同控制框架。该框架以稳定投影双线性 Koopman 升维模型为预测主干，在车辆 Frenet 动力学、团队中心模型、角点刚体映射和连接一致性约束的基础上，引入在线故障诊断与执行器效率辨识、通信质量感知一致性、时延补偿、切换容错控制、阶段-角色调度器和 Lyapunov 型证书记录。与 Adaptive Koopman Embedding baseline 相比，本文不只关注单体系统的模型自适应鲁棒性，而是进一步处理多车载荷运输中的通信失配、执行亏损和团队等效力闭合问题。最新 TF14 图包表明，所提方法在高通信退化和混合故障/噪声场景下能够保持完整路径跟踪，并通过 FDI/FTC 与安全证书把未知故障恢复过程转化为可检测、可重构、可约束的闭环行为。当前图件仍属于单 seed gap-closure 阶段，最终投稿前需补齐 n>=20 多随机种子统计。",
    )
    set_paragraph_text(doc.paragraphs[5], "Koopman 学习；协同运输；模型预测控制；在线故障诊断；切换容错；通信韧性")

    set_paragraph_text(
        doc.paragraphs[12],
        "基于上述分析，本文的比较逻辑不是与历史版本互相比较，而是以 task-aligned Adaptive Koopman Embedding baseline 为主 baseline，并说明 TF14 相对于原始 AKE 思路在四车载荷运输、通信退化、未知执行器故障和稳定性证书方面的扩展。最新 TF14 主方法已经把在线 FDI、执行器效率辨识、切换 FTC、phase-role 调度、通信补偿和证书记录纳入闭环，因此本文贡献必须围绕“稳定投影 Koopman 预测主干 + 网络韧性协同 + 未知故障自治容错 + 证书化安全诊断”组织。",
    )
    set_paragraph_text(
        doc.paragraphs[13],
        "在 Muhammed 等关于受约束环境下多机器人物体运输 MPC 与实时去中心化协同搬运验证，以及 Kennel-Maushart 与 Coros 关于载荷感知轨迹优化和倾覆规避约束的基础上，建立面向四车刚性载荷搬运的载荷中心 Frenet 团队模型，将单车动力学、角点映射、柔性连接误差、团队控制冗余、载荷受力指标和角点法向载荷统一到同一表述中，为未知单车故障如何转化为团队可补偿亏损提供状态基础。",
    )
    set_paragraph_text(
        doc.paragraphs[14],
        "在 Singh 等 adaptive Koopman embedding、Zhao 等 deep bilinear Koopman MPC 以及 Abtahi 等 Frenet 车辆 Koopman 控制的基础上，构建 24 维可学习观测、31 维 lifted state、4 层宽度 128 的 gelu 编码器和稳定投影双线性 Koopman 主干，并使用 ridge_lambda=5e-5、linear_fit_blend=0.60 和 previous_model_blend=0.20 完成离线双线性矩阵拟合，使预测模型兼顾局部拟合能力和谱半径稳定性。",
    )
    set_paragraph_text(
        doc.paragraphs[15],
        "在时延补偿 distributed/cloud MPC、事件触发 DMPC、data-driven predictive consensus control 以及执行器退化容错控制研究的基础上，设计在线 FDI、执行器效率辨识、通信质量感知一致性、时延前推预测、退化安全回退和 reconfigured FTC 的协同控制栈。其核心思想是：先识别故障类型和执行效率，再根据置信度、通信质量和安全裕度选择 nominal、reconfigured 或 safe degraded 模式，从而把不可控的局部执行损失转化为可解释的团队级补偿。",
    )
    set_paragraph_text(
        doc.paragraphs[16],
        "在 Singh 等关于 adaptive Koopman robustness 的对比评估思路与协同运输实验验证组织方式基础上，构建覆盖离线数据、动力学学习、团队轨迹、单车误差、连接一致性、载荷受力、在线诊断、容错切换、稳定性证书和力学安全权衡的 TF14 仿真证据链。当前图包能够作为阶段性机制证据，最终稿仍需用 n>=20 多 seed 统计替换单 seed 图件。"
    )

    replacements = {
        "tf13_pre.ipynb": "tf14_pre",
        "tf13_pre": "tf14_pre",
        "最新版 tf13": "最新 TF14",
        "tf13 主": "TF14 主",
        "tf13 的": "TF14 的",
        "tf13 在": "TF14 在",
        "tf13 已": "TF14 已",
        "tf13 相比": "TF14 相比",
        "tf13": "TF14",
        "TF13": "TF14",
        "tf12": "TF12",
        "上一版按 TF12": "上一版按 TF12/TF13",
        "24 维可学习观测、31 维 lifted state": "24 维可学习观测、31 维 lifted state",
    }
    replace_text_everywhere(doc, replacements)

    # Tighten method text around online adaptation to reflect the TF14 realtime preset.
    set_paragraph_text(
        doc.paragraphs[90],
        "TF14 仍保留在线双线性自适应接口：在线阶段不改变升维映射，只在高置信窗口内修正双线性动力学矩阵。需要区分的是，实时主 preset 为保证控制预算，默认采用缓存/稳定化模型并关闭在线更新；若论文中讨论在线自适应收益，应在非实时或单独消融 preset 中开启，不能把实时主实验和在线更新收益混写。",
    )
    set_paragraph_text(
        doc.paragraphs[106],
        "TF14 主实时配置采用较短预测时域与预算感知调度，典型设置为 horizon=12、max_sqp_iters=1、mpc_decimation_steps=5，并在预算不足时允许 leader-first 或部分车辆跳步求解；error-match/all-solve preset 则提高求解频率以展示精度-计算时间前沿。这样的取舍是为了把在线 FDI、切换 FTC、通信补偿和安全证书纳入同一实时闭环。"
    )
    set_paragraph_text(
        doc.paragraphs[153],
        "4.5 通信退化回退、在线 FDI 与执行器故障容错",
    )
    set_paragraph_text(
        doc.paragraphs[154],
        "TF14 不再假设故障车辆和故障幅值完全已知，而是通过 OnlineFDIMonitorTF14 比较一阶预测残差与实际状态，并结合执行器效率估计区分 residual_only_alert、accel_degraded、steer_degraded、dual_degraded 和 severe_dual_loss。通信质量低时，残差阈值会保守调整，以避免把网络抖动误判为执行器故障。"
    )
    set_paragraph_text(
        doc.paragraphs[164],
        "当 FDI 置信度超过阈值后，ReconfigurableFTCControllerTF14 在 nominal_koopman_mpc、reconfigured_ftc_mpc 和 safe_degraded_consensus 三类模式之间切换。reconfigured 模式先对故障车进行效率逆补偿，再把剩余等效力/力矩缺口分配给健康车辆；severe 或通信质量过低时，控制器进入低激进度安全一致性模式。"
    )
    set_paragraph_text(
        doc.paragraphs[172],
        "如果把 AKE-baseline 概括为“离线 Koopman 主干 + 在线自适应 + MPC”，那么 TF14 的主方法增量是沿预测、通信、执行和证书四条链路闭合失效源。预测层用稳定投影双线性 Koopman 抑制长时域外推风险；通信层用质量感知一致性和时延补偿减少协调偏差；执行层用在线 FDI 与切换 FTC 处理未知故障；证书层用 Lyapunov 型 margin 和 practical-bound 记录把安全假设落到闭环数据。"
    )
    set_paragraph_text(
        doc.paragraphs[173],
        "这四条链路并不是彼此独立的。更准确地说，Koopman 主干减少“送入 MPC 的名义预测偏差”，通信补偿减少“从 MPC 输出到团队执行之间的协调偏差”，FTC 减少“从团队命令到实际载荷受力之间的执行偏差”，证书记录则监测“切换控制是否仍处在可证明的实用稳定边界内”。"
    )
    set_paragraph_text(
        doc.paragraphs[174],
        "从审稿视角看，TF14 与 AKE-baseline 最本质的区别在于：AKE-baseline 主要回答模型失配后的 Koopman 控制鲁棒性，而 TF14 进一步回答当这种 Koopman 控制进入通信退化、执行器未知故障和刚性载荷耦合的多车系统时，怎样把模型层优势转化为团队级任务完成能力与可诊断安全性。"
    )

    # Insert method addendum before section 5.
    sec5_anchor = next(p for p in doc.paragraphs if p.text.strip().startswith("5 仿真实验"))
    method_blocks = build_method_update_blocks(doc)
    move_blocks_before(sec5_anchor, method_blocks)

    # Remove legacy TF13 transition experiment sections and insert latest TF14 evidence.
    conclusion_anchor = next(p for p in doc.paragraphs if p.text.strip().startswith("6 结论"))
    legacy_start = next((p for p in doc.paragraphs if p.text.strip().startswith("5.6 复杂路径")), None)
    if legacy_start is not None:
        remove_blocks_between(legacy_start, conclusion_anchor)
    exp_blocks = build_experiment_update_blocks(doc)
    move_blocks_before(conclusion_anchor, exp_blocks)

    # Rewrite conclusion in TF14 terms after insertions.
    # Locate conclusion again because paragraph list was extended.
    paragraphs = doc.paragraphs
    conclusion_idx = next(i for i, p in enumerate(paragraphs) if p.text.strip().startswith("6 结论"))
    set_paragraph_text(
        paragraphs[conclusion_idx + 1],
        "本文围绕四车刚性载荷协同运输任务，在已有文献综述稿的基础上进一步纳入 TF14 方法更新，形成了从系统建模、稳定投影双线性 Koopman 升维、通信质量感知一致性、时延补偿、在线 FDI、执行器效率辨识、切换 FTC、阶段-角色调度和 Lyapunov 型证书记录到载荷安全诊断的完整控制框架。",
    )
    set_paragraph_text(
        paragraphs[conclusion_idx + 2],
        "与 task-aligned Adaptive Koopman baseline 相比，TF14 的核心增量不再只是更强的名义 Koopman 主干，而是把模型预测、网络通信、执行器故障和稳定性证书四条脆弱链路同时纳入一个可解释闭环。最新图包显示，TF14 在高通信退化和混合故障/噪声场景下能够保持完整路径推进，并通过 FDI/FTC 将未知故障恢复过程转化为可检测、可重构和可约束的团队行为。",
    )
    set_paragraph_text(
        paragraphs[conclusion_idx + 3],
        "当前稿件仍需谨慎处理统计边界：E1/E2/E3/E7 已形成阶段性证据闭环，但 minimum_group_n=1，不能替代最终多 seed 显著性结论。下一步应优先补齐 n>=20 多随机种子统计、真实 Koopman DNN 多种子重训练、FDI 分类准确率、severe fault fallback、phase-role 高曲率收益、runtime profiling 以及 matched TF13/TF14 精度-实时性对比，使该方法更接近一区 top 期刊对完整性、可解释性和可信度的要求。",
    )

    normalize_sections(doc)
    normalize_source_tables(doc)
    coerce_docx_width_numbers(doc)
    doc.save(str(OUT_DOCX))
    print(OUT_DOCX)
    print("paragraphs", len(doc.paragraphs), "tables", len(doc.tables), "inline_shapes", len(doc.inline_shapes))


if __name__ == "__main__":
    main()
