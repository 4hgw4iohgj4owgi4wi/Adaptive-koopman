from __future__ import annotations

import json
import sys
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"D:\LEARNING\ZNN\ZNN\Adaptive-koopman\Adaptive-koopman-main")
SRC = ROOT / "paper_dcn_tf12_draft" / "tf14_method_update_zh_with_evidence_figures_2026-05-09.docx"
OUT = ROOT / "paper_dcn_tf12_draft" / "tf14_method_update_zh_with_remaining_figures_2026-05-09.docx"
FIG_DIR = ROOT / "tf14_remaining_experiments_20260509" / "figures"
DATA_DIR = ROOT / "tf14_remaining_experiments_20260509" / "data"

DOC_SKILL = Path(
    r"C:\Users\lj\.codex\plugins\cache\openai-primary-runtime\documents\26.430.10722\skills\documents\scripts"
)
sys.path.append(str(DOC_SKILL))
from table_geometry import apply_table_geometry, column_widths_from_weights, section_content_width_dxa  # noqa: E402


BLUE = "0E5A78"
LIGHT_BLUE = "DDEFF5"
PALE_YELLOW = "FFF2CC"
PALE_GREEN = "E2F0D9"


def set_font(run, size: float | None = None, bold: bool | None = None, color: str | None = None) -> None:
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def add_heading(doc: Document, text: str, level: int = 1):
    p = doc.add_heading(text, level=level)
    for r in p.runs:
        set_font(r, color=BLUE)
    return p


def add_para(doc: Document, text: str, size: float = 10.5):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    set_font(r, size=size)
    return p


def add_callout(doc: Document, title: str, body: str, fill: str = PALE_YELLOW) -> None:
    width = section_content_width_dxa(doc.sections[-1])
    table = doc.add_table(rows=1, cols=1)
    apply_table_geometry(table, [width], table_width_dxa=width, indent_dxa=0)
    cell = table.cell(0, 0)
    shade(cell, fill)
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(title)
    set_font(r, size=10.5, bold=True, color=BLUE)
    p2 = cell.add_paragraph()
    r2 = p2.add_run(body)
    set_font(r2, size=10.0)
    doc.add_paragraph()


def set_cell(cell, text: str, header: bool = False) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(str(text))
    set_font(r, size=8.0, bold=header, color=BLUE if header else None)
    if header:
        shade(cell, LIGHT_BLUE)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], weights: list[float]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    for i, h in enumerate(headers):
        set_cell(table.cell(0, i), h, header=True)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell(cells[i], value, header=False)
    widths = column_widths_from_weights(weights, section_content_width_dxa(doc.sections[-1]))
    apply_table_geometry(table, widths, table_width_dxa=sum(widths), indent_dxa=0)
    table.style = "Table Grid"
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)
    doc.add_paragraph()


def add_figure(doc: Document, filename: str, caption: str, width_in: float = 6.25) -> None:
    path = FIG_DIR / filename
    if not path.exists():
        add_para(doc, f"[缺失图件：{filename}]", size=9.0)
        return
    pic = doc.add_picture(str(path), width=Inches(width_in))
    pic_par = doc.paragraphs[-1]
    pic_par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run(caption)
    set_font(r, size=8.5, color="666666")


def replace_validation_paragraph(doc: Document) -> None:
    old = "下表检查 TF14 的每个主要改进是否已有对应实验验证。"
    new_text = (
        "下表检查 TF14 的每个主要改进是否已有对应实验验证。结合最新 "
        "tf14_remaining_experiments_20260509 图包，控制效果、真实逐模块消融、"
        "Lyapunov 型证书时间线以及载荷/连接安全已经具备单 seed 阶段性证据；"
        "但这些证据仍不能作为最终统计结论，因为当前 minimum_group_n=1、"
        "final_statistics_ready=false。Koopman DNN 预测验证、通信模块消融、FDI 分类、"
        "FTC severe fallback、phase-role 高曲率收益、runtime profiling 和 matched TF13 对比仍需补图。"
    )
    for p in doc.paragraphs:
        if old in p.text:
            p.text = ""
            r = p.add_run(new_text)
            set_font(r, size=10.5)
            break


def main() -> None:
    manifest = json.loads((DATA_DIR / "tf14_remaining_manifest.json").read_text(encoding="utf-8"))
    summary = pd.read_csv(DATA_DIR / "tf14_remaining_summary.csv")
    doc = Document(SRC)
    replace_validation_paragraph(doc)

    add_heading(doc, "16. 最新 remaining 图包的入文更新", 1)
    add_para(
        doc,
        "本节纳入 tf14_remaining_experiments_20260509/figures 中可作为阶段性证据的图件。"
        "该图包较上一版新增了主对比时间序列、真实模块消融、稳定性证书时间线以及载荷/连接安全对比，"
        "因此可以补强实验章节中“控制效果、模块贡献、稳定性机制、安全约束”的叙事链条。"
    )
    add_callout(
        doc,
        "写作口径",
        f"当前 manifest 显示 experiments_present={manifest.get('experiments_present')}，"
        f"scenarios_present={manifest.get('scenarios_present')}，minimum_group_n={manifest.get('minimum_group_n')}，"
        f"final_statistics_ready={manifest.get('final_statistics_ready')}。因此以下图件应写作“preliminary single-seed evidence / 阶段性代表性结果”，"
        "不能写作最终 n>=20 多 seed 统计。终稿中需要用相同图型替换为 n>=20 的均值、置信区间和显著性结果。"
    )

    included_rows = [
        [
            "R0_error_timeseries",
            "主文候选/阶段性",
            "展示 high-communication degradation 与 mixed fault/noise 下 e_y、e_s 时间序列，说明 TF14 没有出现发散。",
            "标题和正文必须说明 n=1；后续需要 mean±95%CI 阴影和故障/通信窗口标注。",
        ],
        [
            "R0_main_comparison_summary",
            "阶段性主对比",
            "给出 RMSE_y、max |e_y|、full path rate 的两场景汇总。",
            "不能作为最终主结果；max |e_y| 对 TF14 不完全有利，应如实解释为连接安全优先带来的峰值代价。",
        ],
        [
            "R2 heatmap + delta",
            "消融候选",
            "证明 no_PPC_progress_guard、no_comm_aware、no_FDI/no_FTC 等模块退化后出现 RMSE、连接或证书损失。",
            "色标需在终稿中统一为 positive=worse；当前仍需 n>=20。",
        ],
        [
            "R3 certificate timelines",
            "稳定性机制候选",
            "显示 V、upper bound、margin、certificate_ok 和模式切换，可支撑 practical stability 证书机制。",
            "应写成 practical/ISS stability 证据，不写成严格逐步收缩；后续补多 seed 统计。",
        ],
        [
            "R7 payload/connection safety",
            "安全权衡候选",
            "显示连接利用率低于阈值且 violation count 为 0，支撑连接安全约束有效。",
            "同时显示 force peak/corner load spread 上升，必须写成安全-受力代价权衡，而不是全面优于 baseline。",
        ],
    ]
    add_table(doc, ["图组", "建议位置", "可支撑内容", "入文限制"], included_rows, [1.5, 1.3, 3.4, 3.4])

    add_heading(doc, "16.1 主对比与误差时间序列", 2)
    add_para(
        doc,
        "R0 图组用于补强系统级控制效果。当前两类压力场景中，TF14 main、phase-role 和 error-match 均保持完整路径完成；"
        "横向 RMSE 和连接安全指标呈改善趋势。但由于单 seed 结果中 max |e_y| 并非全项占优，论文描述应强调“跟踪稳定性与连接安全改善”，"
        "而非简单写成所有误差指标全面降低。"
    )
    add_figure(
        doc,
        "R0_error_timeseries.png",
        "图 16-1  高通信退化和混合故障/噪声下的横向与纵向误差时间序列。当前为单 seed 代表性结果，后续需替换为多 seed 均值与置信区间。",
        6.15,
    )
    add_figure(
        doc,
        "R0_main_comparison_summary.png",
        "图 16-2  两类压力场景下的主指标汇总。该图可作为阶段性主对比，终稿需补 n>=20 误差条。",
        6.25,
    )

    add_heading(doc, "16.2 真实逐模块消融", 2)
    add_para(
        doc,
        "R2 图组比旧版 available variants 更接近审稿人要求的真实消融：每次只关闭一个模块，并观察 RMSE、连接 RMS、force peak、certificate ok 和 step p95。"
        "当前结果显示，PPC/progress guard 是避免纵向进度崩溃和证书失效的关键；communication-aware 机制对连接误差和证书裕度影响明显；"
        "FDI/FTC 关闭后在 mixed fault/noise 场景中会放大横向误差和峰值误差。"
    )
    add_figure(
        doc,
        "R2_module_ablation_heatmap_dlc_comm_noise_high.png",
        "图 16-3  高通信退化场景下的多指标逐模块消融热力图。色块表示相对 full TF14 的变化，终稿需统一“正值/负值”的好坏方向。",
        6.0,
    )
    add_figure(
        doc,
        "R2_module_ablation_heatmap_dlc_mixed_fault_noise.png",
        "图 16-4  混合故障/噪声场景下的多指标逐模块消融热力图。no_PPC_progress_guard 的退化最明显。",
        6.0,
    )
    add_figure(
        doc,
        "R2_true_module_ablation_delta.png",
        "图 16-5  两类压力场景下各消融项相对 full TF14 的横向 RMSE 增量。该图适合作为消融摘要图。",
        6.15,
    )

    add_heading(doc, "16.3 Lyapunov 型证书与模式切换", 2)
    add_para(
        doc,
        "R3 时间线图是当前最适合放入稳定性证明对应实验的位置。它把 Lyapunov 型函数 V、预测上界、收缩裕度、certificate_ok 以及控制模式放在同一时间轴上，"
        "能够说明 TF14 的稳定性模块不是纯理论附会，而是在闭环运行中持续监测。mixed fault/noise 图中出现 nominal 到 reconfigured 的模式转移，更适合解释 FDI/FTC 与证书之间的联动。"
    )
    add_figure(
        doc,
        "R3_certificate_timeline_dlc_comm_noise_high.png",
        "图 16-6  高通信退化场景下的 practical-stability certificate 时间线。",
        6.1,
    )
    add_figure(
        doc,
        "R3_certificate_timeline_dlc_mixed_fault_noise.png",
        "图 16-7  混合故障/噪声场景下的证书时间线与模式切换。该图可用于支撑 reconfigured FTC 的稳定性机制。",
        6.1,
    )

    add_heading(doc, "16.4 载荷与连接安全", 2)
    add_para(
        doc,
        "R7 图组用于修正此前只看跟踪误差的叙事。TF14 的主要优势不是单纯降低所有误差，而是在通信退化和故障条件下把连接利用率维持在阈值内，并保持 connection violation count 为 0。"
        "同时，载荷 force peak 和 corner load spread 有上升趋势，说明 TF14 更像是通过主动补偿和载荷重分配换取连接安全。这一点需要在论文中写成“安全约束收益与力学代价的权衡”。"
    )
    add_figure(
        doc,
        "R7_payload_connection_safety_dlc_comm_noise_high.png",
        "图 16-8  高通信退化场景下的载荷受力、连接利用率、违反次数和角点载荷扩展。",
        6.1,
    )
    add_figure(
        doc,
        "R7_payload_connection_safety_dlc_mixed_fault_noise.png",
        "图 16-9  混合故障/噪声场景下的载荷与连接安全指标。TF14 保持零连接违反，但存在更高受力代价。",
        6.1,
    )

    add_heading(doc, "16.5 更新后的实验缺口", 2)
    add_para(
        doc,
        "纳入本图包后，E0/E2/E3/E7 的实验链条已经有阶段性图件支撑；但终稿仍必须补齐统计强度和未覆盖模块。"
    )
    gap_rows = [
        ["P0", "E0 主对比", "将 R0 图组扩展为 n>=20，多 seed 均值/95%CI、boxplot、failure count 和 solver success。"],
        ["P0", "E2 真消融", "所有消融至少 n>=20；热力图色标统一为 positive=worse，并补统计显著性。"],
        ["P0", "E3 证书", "保留 R3 时间线作为代表案例，同时补 mode-wise margin 分布和 ultimate bound 统计。"],
        ["P0", "E7 载荷安全", "补 force rate/jerk、corner load spread 时间线，并解释 force peak 上升的安全代价。"],
        ["P0", "E1 Koopman DNN", "补 one-step、multi-step rollout、spectral radius/stable projection，证明新 Koopman 网络本身有效。"],
        ["P1", "E4/E5/E6", "补通信模块消融、FDI 分类混淆矩阵/ROC、FTC severe fallback。"],
        ["P1", "E8/E9/E10", "补 phase-role 高曲率收益、runtime profiling、matched TF13 对比。"],
    ]
    add_table(doc, ["优先级", "缺口", "需要补的内容"], gap_rows, [0.8, 1.8, 6.4])
    add_callout(
        doc,
        "本节结论",
        "最新 remaining 图包可以显著增强中文稿的实验叙事，但只能作为阶段性图件入文。正式投稿前，需要用同样版式替换为 n>=20 多 seed 结果，并补齐 Koopman DNN、通信、FDI、FTC severe、phase-role 高曲率、runtime 和 TF13 matched 对比。",
        PALE_GREEN,
    )

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
