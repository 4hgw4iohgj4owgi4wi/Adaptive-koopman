from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


ROOT = Path(r"D:\LEARNING\ZNN\ZNN\Adaptive-koopman\Adaptive-koopman-main")
DRAFT = ROOT / "paper_dcn_tf12_draft"
SRC_DOCX = DRAFT / "tf14_method_update_zh_with_gap_closure_figures_equations_2026-05-09.docx"
OUT_DOCX = DRAFT / "tf14_method_update_zh_with_gap_closure_quantitative_review_2026-05-09.docx"
CSV_DIR = ROOT / "tf14_final_gap_closure_20260509" / "source"


def set_run_font(run, size=None, bold=None, color=None, name="宋体"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_twips):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_twips))
    tc_w.set(qn("w:type"), "dxa")


def set_table_width(table, widths):
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        table._tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)


def set_cell_text(cell, text, bold=False, fill=None, color=(34, 34, 34), align="left"):
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    if fill:
        shade_cell(cell, fill)
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.06
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(str(text))
    set_run_font(run, size=8.7, bold=bold, color=color, name="宋体")


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(run, name="黑体")
    return p


def add_body(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(24)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    set_run_font(run, size=10.5)
    return p


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for cell, header in zip(table.rows[0].cells, headers):
        set_cell_text(cell, header, bold=True, fill="EAF2F8", color=(31, 78, 121), align="center")
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            set_cell_text(cell, value)
    set_table_width(table, widths)
    doc.add_paragraph()
    return table


def replace_text_in_doc(doc, replacements):
    for paragraph in doc.paragraphs:
        for old, new in replacements.items():
            if old in paragraph.text:
                paragraph.text = paragraph.text.replace(old, new)
                for run in paragraph.runs:
                    set_run_font(run, size=9 if paragraph.text.startswith("图 17") else 10.5)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for old, new in replacements.items():
                        if old in paragraph.text:
                            paragraph.text = paragraph.text.replace(old, new)
                            for run in paragraph.runs:
                                set_run_font(run, size=8.7)


def fmt(x, nd=3):
    return f"{x:.{nd}f}"


def pct(x):
    return f"{x * 100:+.1f}%"


def build_quant_rows():
    e1 = pd.read_csv(CSV_DIR / "E1_koopman_prediction_validation.csv")
    one = e1.dropna(subset=["one_step_rmse"]).groupby("model")["one_step_rmse"].mean()
    roll = e1.dropna(subset=["horizon"]).copy()
    roll["horizon"] = roll["horizon"].astype(int)
    h60 = roll[roll["horizon"] == 60].set_index("model")["rollout_mean_l2"]
    spec = e1.dropna(subset=["spectral_radius_before", "spectral_radius_after"]).iloc[0]
    e1_rows = [
        (
            "一阶预测 RMSE 均值",
            f"linear={fmt(one['linear'])}，bilinear={fmt(one['bilinear'])}，stable bilinear={fmt(one['stable_bilinear'])}",
            "bilinear 在一阶平均误差上略优，可说明双线性项有局部拟合收益。",
            "不能写成 stable bilinear 全面降低预测误差；稳定化会引入轻微拟合代价。",
        ),
        (
            "60-step 开环 rollout",
            f"linear={fmt(h60['linear'],2)}，bilinear={fmt(h60['bilinear'],2)}，stable bilinear={fmt(h60['stable_bilinear'],2)}",
            "长时开环预测暴露出模型外推风险，因此论文应强调闭环 MPC、稳定投影和约束保护的必要性。",
            "不能把 E1 当作“长时预测显著优于 baseline”的证据。",
        ),
        (
            "稳定投影谱半径",
            f"投影前={fmt(spec['spectral_radius_before'],4)}，投影后={fmt(spec['spectral_radius_after'],4)}",
            "可作为稳定化网络结构的直接证据：谱半径被压到 1 以下。",
            "仍需配合 closed-loop 统计证明控制收益，而不是只看线性算子谱半径。",
        ),
    ]

    e2 = pd.read_csv(CSV_DIR / "E2_positive_worse_ablation.csv")
    top_rows = []
    for scenario, zh in [
        ("dlc_comm_noise_high", "高通信退化"),
        ("dlc_mixed_fault_noise", "混合故障/噪声"),
    ]:
        sub = e2[(e2["scenario"] == scenario) & (e2["positive_worse_delta"] > 0)]
        top = sub.sort_values("positive_worse_delta", ascending=False).head(4)
        desc = "；".join(
            f"{r.method}/{r.metric} +{r.positive_worse_delta:.3g}" for r in top.itertuples()
        )
        top_rows.append(
            (
                zh,
                desc,
                "热力图可用于说明哪些模块被移除后最容易退化。",
                "当前 num_runs=1，只能写趋势和诊断，不能写显著性。",
            )
        )

    e3 = pd.read_csv(CSV_DIR / "E3_certificate_mode_margin_stats.csv")
    e3_rows = []
    for r in e3.itertuples():
        e3_rows.append(
            (
                r.scenario.replace("dlc_", ""),
                r.mode,
                f"ok={fmt(r.ok_ratio,1)}，min margin={fmt(r.margin_min,4)}，p05={fmt(r.margin_p05,4)}",
                f"sqrt upper p95={fmt(r.ultimate_bound_sqrt_upper_p95,3)}",
            )
        )

    e7 = pd.read_csv(CSV_DIR / "E7_force_rate_jerk_corner_summary.csv")
    e7_rows = []
    for scenario, zh in [
        ("dlc_comm_noise_high", "高通信退化"),
        ("dlc_mixed_fault_noise", "混合故障/噪声"),
    ]:
        sub = e7[e7["scenario"] == scenario].set_index("method")
        base = sub.loc["baseline"]
        tf = sub.loc["tf14_main"]
        noftc = sub.loc["no_ftc_switching"]
        vs_base = (
            f"force peak {pct(tf.force_norm_peak/base.force_norm_peak-1)}，"
            f"rate {pct(tf.force_norm_rate_peak/base.force_norm_rate_peak-1)}，"
            f"jerk {pct(tf.force_norm_jerk_peak/base.force_norm_jerk_peak-1)}"
        )
        vs_noftc = (
            f"force peak {pct(tf.force_norm_peak/noftc.force_norm_peak-1)}，"
            f"rate {pct(tf.force_norm_rate_peak/noftc.force_norm_rate_peak-1)}，"
            f"jerk {pct(tf.force_norm_jerk_peak/noftc.force_norm_jerk_peak-1)}，"
            f"corner peak {pct(tf.corner_load_spread_peak/noftc.corner_load_spread_peak-1)}"
        )
        e7_rows.append((zh, vs_base, vs_noftc, "应写成恢复能力与力学代价的权衡，而不是“全部力学指标优于 baseline”。"))

    return e1_rows, top_rows, e3_rows, e7_rows


def main():
    doc = Document(str(SRC_DOCX))
    replacements = {
        "证明新 Koopman 网络具有更低预测误差和稳定投影效果；但还不是多种子 DNN 重训练。": "说明双线性 Koopman 在一阶预测上有局部收益，并暴露长时开环 rollout 风险；稳定投影可把谱半径压到 1 以下，但还不是多种子 DNN 重训练证据。",
        "该图可用于说明新 Koopman 网络在离线缓存筛选中的建模优势，但不能替代真正的多随机种子 DNN 重训练统计。": "该图用于说明双线性项的一阶拟合收益、长时开环外推风险，以及稳定投影的谱半径约束效果；不能替代真正的多随机种子 DNN 重训练统计。",
        "E1 来自缓存离线筛选数据，不等价于完整 DNN 多种子重训练；": "E1 来自缓存离线筛选数据，不等价于完整 DNN 多种子重训练；并且它更适合作为“稳定投影与闭环控制必要性”的证据，而不是“长时开环预测全面胜出”的证据；",
    }
    replace_text_in_doc(doc, replacements)

    e1_rows, e2_rows, e3_rows, e7_rows = build_quant_rows()

    add_heading(doc, "18. gap-closure 图包的定量解读与审稿口径修正", level=1)
    add_body(
        doc,
        "继续检查 source 数据后，需要对第 17 节的叙述做更精确的收束：这批图可以补强论文证据链，但不能被写成所有模块都已经完成最终统计验证。尤其是 E1 与 E7 要谨慎。E1 的价值不是证明长时开环预测全面更好，而是证明双线性局部拟合、稳定投影和闭环 MPC 保护都必要；E7 的价值也不是证明 TF14 的力峰值总是更低，而是把故障恢复速度、连接误差、安全约束和力学代价放在同一个诊断框架里。",
    )

    add_heading(doc, "18.1 E1 Koopman 学习结果的正确写法", level=2)
    add_table(
        doc,
        ["指标", "当前数值", "可以支持的结论", "不能这样写"],
        e1_rows,
        [1900, 2600, 3000, 1860],
    )

    add_heading(doc, "18.2 E2 消融热力图的主要退化来源", level=2)
    add_table(
        doc,
        ["场景", "top positive=worse 项", "可以支持的结论", "边界"],
        e2_rows,
        [1500, 3550, 2650, 1660],
    )
    add_body(
        doc,
        "positive=worse 的最大好处是审稿人不需要再反向解释颜色：红色越深表示去掉该模块后相对 TF14 主方法越差。当前最明显的退化集中在 no_comm_aware、no_ppc_progress_guard，以及混合故障下的 no_ftc_switching/FDI 相关项，这与本文强调的通信鲁棒、进度保护和容错切换模块是对应的。",
    )

    add_heading(doc, "18.3 E3 稳定性证书与 E7 力学安全的合并解释", level=2)
    add_table(
        doc,
        ["场景", "控制模态", "证书裕度", "practical-bound proxy"],
        e3_rows,
        [1800, 2500, 3100, 1960],
    )
    add_table(
        doc,
        ["场景", "TF14 main 相对 baseline", "TF14 main 相对 no FTC switching", "论文表述建议"],
        e7_rows,
        [1500, 2600, 2850, 2410],
    )
    add_body(
        doc,
        "E3 的证书统计给出一个较稳的写法：当前三个模态/场景的 ok ratio 均为 1.0，最小裕度仍为正，因此可作为 Lyapunov/ISS 证明的仿真闭环证据。E7 则要更克制：相对 baseline，TF14 在强扰动或故障恢复阶段可能付出更高 force-rate/jerk 代价；但相对 no_ftc_switching，混合故障场景下 TF14 main 显著降低了力峰值、力变化率、jerk 和角点载荷峰值。因此更合理的创新叙述是：TF14 通过 FTC 和调度保护把不可控的故障冲击转化为可解释、可约束、可诊断的恢复过程。",
    )

    add_heading(doc, "18.4 后续实验优先级", level=2)
    priority_rows = [
        ("P0", "n≥20 多随机种子统计", "对 E2/E3/E7 统一补均值、标准差、置信区间与显著性检验；这是最终投稿版必须补齐的统计可信度。"),
        ("P0", "真实 Koopman DNN 多种子重训练", "E1 当前是 cached offline screening；最终需要 linear/bilinear/stable-bilinear 在相同数据划分下重新训练并统计。"),
        ("P1", "E7 平顺性代价解释", "补充约束激活率、恢复时间、连接误差峰值和力学指标之间的 trade-off 图，避免被审稿人误读为力学指标退化。"),
        ("P1", "FDI/FTC severe 专项", "补 FDI 检测延迟、分类准确率、误报漏报，以及 severe fault fallback 的安全包络结果。"),
        ("P1", "TF13 matched 与 runtime", "补同场景同扰动对齐的 TF13 matched 对比和求解时间分布，回应方法复杂度与实时性问题。"),
    ]
    add_table(
        doc,
        ["优先级", "实验/图件", "为什么必须补"],
        priority_rows,
        [1000, 2600, 5760],
    )
    add_body(
        doc,
        "以上修正后，中文稿的证据链更接近顶刊审稿口径：不把阶段性图包过度包装为最终结论，而是明确区分“已经能证明的机制”“只能作为趋势的诊断图”和“最终投稿前必须补齐的统计实验”。",
    )

    doc.save(str(OUT_DOCX))
    print(OUT_DOCX)


if __name__ == "__main__":
    main()
