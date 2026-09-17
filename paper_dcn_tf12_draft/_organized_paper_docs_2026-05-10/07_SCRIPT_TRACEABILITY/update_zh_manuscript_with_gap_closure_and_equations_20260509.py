from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"D:\LEARNING\ZNN\ZNN\Adaptive-koopman\Adaptive-koopman-main")
DRAFT = ROOT / "paper_dcn_tf12_draft"
SRC_DOCX = DRAFT / "tf14_method_update_zh_with_remaining_figures_2026-05-09.docx"
OUT_DOCX = DRAFT / "tf14_method_update_zh_with_gap_closure_figures_equations_2026-05-09.docx"
EQ_DIR = DRAFT / "equation_images_20260509"
GAP_FIG_DIR = ROOT / "tf14_final_gap_closure_20260509" / "figures"


FORMULA_REPLACEMENTS = {
    "dot X_i = v_x,i cos psi_i - v_y,i sin psi_i,  dot Y_i = v_x,i sin psi_i + v_y,i cos psi_i,  dot psi_i = r_i.": "eq_global_kinematics.png",
    "dot s_i = (v_x,i cos e_psi,i - v_y,i sin e_psi,i)/(1 - kappa(s_i)e_y,i)": "eq_frenet_s.png",
    "dot e_y,i = v_x,i sin e_psi,i + v_y,i cos e_psi,i,  dot e_psi,i = r_i - kappa(s_i) dot s_i": "eq_frenet_error.png",
    "dot v_y,i = -(2C_f+2C_r)/(m v_x,i) v_y,i + (-v_x,i -(2C_f l_f-2C_r l_r)/(m v_x,i)) r_i + (2C_f/m) delta_i": "eq_bicycle_vy.png",
    "dot r_i = -(2C_f l_f-2C_r l_r)/(I_z v_x,i) v_y,i - (2C_f l_f^2+2C_r l_r^2)/(I_z v_x,i) r_i + (2C_f l_f/I_z) delta_i": "eq_bicycle_r.png",
    "dot v_x,i = a_x,i + r_i v_y,i,  v_x,i >= 0.5 m/s.": "eq_bicycle_vx.png",
    "z_{k+1} = A z_k + B u_k + sum_{j=1}^{2} u_{k,j} N_j z_k + epsilon_k.": "eq_koopman_bilinear.png",
    "hat x_{i,k|k-1}=f_i(x_{i,k-1},u^{cmd}_{i,k-1}),  r_{i,k}=x_{i,k}-hat x_{i,k|k-1}": "eq_fdi_residual.png",
    "rho_{i,k}=sqrt(sum_l (r_{i,k,l}/eta_l)^2),  EWMA_k=beta EWMA_{k-1}+(1-beta)rho_{i,k}": "eq_fdi_ewma.png",
    "u^{act}_{i,k}=Gamma_{i,k}u^{cmd}_{i,k},  Gamma_{i,k}=diag(gamma_delta,i,k, gamma_a,i,k).": "eq_actuator_efficiency.png",
    "u_i^R = Pi_U( hat Gamma_i^{-1} u_i^MPC + Delta u_i^{redist} ),  sum_i Delta u_i^{redist} approx missing team effort.": "eq_ftc_reconfig.png",
    "u_i^S = Pi_U( diag(alpha_delta, alpha_a) u_i^MPC + Delta u_i^{safe} ).": "eq_ftc_safe.png",
    "u_i^{out}=Pi_U(u_i^{FTC} + [Delta delta_i^{phase}, Delta a_i^{phase}]^T).": "eq_phase_role_output.png",
    "e_{k+1}=A_{sigma_k}e_k + B_{sigma_k} tilde Gamma_k u_k + d_k,": "eq_switched_error.png",
    "V_{sigma_{k+1}}(e_{k+1}) <= (1 - lambda_sigma Delta t) V_{sigma_k}(e_k) + c_d ||d_k||^2 + c_g ||Gamma_k - hat Gamma_k||^2.": "eq_lyapunov_iss.png",
}


GAP_FIGURES = [
    (
        "E1_koopman_prediction_validation.png",
        "图 17-1  Koopman 预测验证的阶段性闭环证据。左侧是一阶预测误差，中间是多步 rollout 的平均偏差，右侧展示稳定投影前后的谱半径。该图可用于说明新 Koopman 网络在离线缓存筛选中的建模优势，但不能替代真正的多随机种子 DNN 重训练统计。",
    ),
    (
        "E2_positive_worse_heatmap_dlc_comm_noise_high.png",
        "图 17-2  高通信退化场景下的 positive=worse 消融热力图。颜色定义已经改为“正值表示相对 TF14 主方法变差”，比此前 R2 语义更适合论文审稿阅读；当前样本数仍为 n=1。",
    ),
    (
        "E2_positive_worse_heatmap_dlc_mixed_fault_noise.png",
        "图 17-3  混合故障与噪声场景下的 positive=worse 消融热力图。该图能直观说明去掉通信补偿、容错切换或相位角色模块后，多指标退化集中出现的位置；当前仍需最终 n≥20 统计确认。",
    ),
    (
        "E3_certificate_margin_bound_distribution.png",
        "图 17-4  模态证书裕度与 practical-bound 统计。该图用于补强稳定性证明与仿真诊断之间的对应关系，说明切换模态下证书裕度保持为正，并给出实际误差上界的分布范围。",
    ),
    (
        "E7_force_rate_jerk_corner_timeline_dlc_comm_noise_high.png",
        "图 17-5  高通信退化场景下的合力变化率、jerk 与角点载荷时间线。该图适合支撑“TF14 不是只看路径误差，也显式约束运输过程中的平顺性与载荷分配”的论述；若峰值高于 baseline，应解释为快速纠偏与安全裕度之间的权衡。",
    ),
    (
        "E7_force_rate_jerk_corner_timeline_dlc_mixed_fault_noise.png",
        "图 17-6  混合故障与噪声场景下的合力变化率、jerk 与角点载荷时间线。该图可以放入正文或补充材料，用于说明故障切换后系统仍能维持可解释的力学响应；最终稿仍需配合统计表报告均值、标准差与显著性。",
    ),
]


def set_run_font(run, size=None, bold=None, color=None, name="宋体"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def clear_paragraph(paragraph):
    for child in list(paragraph._p):
        if child.tag.endswith("}r") or child.tag.endswith("}hyperlink"):
            paragraph._p.remove(child)


def insert_equation_image(paragraph, image_path, width_inches=5.8):
    clear_paragraph(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(7)
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=Inches(width_inches))


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(run, name="黑体")
    return p


def add_body_paragraph(doc, text, keep_with_next=False):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(24)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.keep_with_next = keep_with_next
    run = p.add_run(text)
    set_run_font(run, size=10.5, name="宋体")
    return p


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text, bold=False, color=(34, 34, 34), fill=None):
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    if fill:
        shade_cell(cell, fill)
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.08
    run = p.add_run(text)
    set_run_font(run, size=9, bold=bold, color=color, name="宋体")


def set_cell_width(cell, width_twips):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_twips))
    tc_w.set(qn("w:type"), "dxa")


def set_table_width(table, widths):
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")

    tbl_grid = table._tbl.tblGrid
    if tbl_grid is None:
        tbl_grid = OxmlElement("w:tblGrid")
        table._tbl.insert(0, tbl_grid)
    for child in list(tbl_grid):
        tbl_grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        tbl_grid.append(col)

    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)


def add_gap_table(doc):
    rows = [
        ("E1", "Koopman 预测验证", "可用作阶段性主文/补充图", "证明新 Koopman 网络具有更低预测误差和稳定投影效果；但还不是多种子 DNN 重训练。"),
        ("E2", "positive=worse 消融热力图", "可用，建议正文放一张、补充材料放一张", "颜色语义清楚，能直观看到去掉模块后的指标退化；目前 n=1，应标注为诊断版。"),
        ("E3", "证书裕度与 practical bound", "可用作稳定性证据补强", "连接 Lyapunov/ISS 证明和仿真统计，说明切换模态裕度为正。"),
        ("E7", "力变化率、jerk、角点载荷", "可用作力学安全性补充图", "支撑平顺性和载荷分配分析；若 TF14 局部峰值更大，应表述为故障恢复速度与力学平顺性的权衡。"),
    ]
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    headers = ["编号", "图件类别", "入稿判断", "审稿口径"]
    widths = [820, 2200, 2300, 4040]
    for cell, header in zip(table.rows[0].cells, headers):
        set_cell_text(cell, header, bold=True, color=(31, 78, 121), fill="EAF2F8")
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            set_cell_text(cells[idx], value)
    set_table_width(table, widths)
    for row in table.rows:
        row.height = None
    doc.add_paragraph()


def add_status_table(doc):
    rows = [
        ("已部分补齐", "E1", "已有缓存离线预测筛选图，可说明结构有效；最终仍需真实多随机种子训练。"),
        ("已部分补齐", "E2", "已把消融热力图改为 positive=worse，读者不再需要反向解释颜色含义；最终仍需 n≥20。"),
        ("已部分补齐", "E3", "已有模态证书裕度和 practical-bound 统计，可支撑稳定性证明的仿真闭环。"),
        ("已部分补齐", "E7", "已有合力变化率、jerk 和角点载荷时间线，可回应运输平顺性与载荷安全问题。"),
        ("仍需补强", "E4/E5/E6/E8/E9/E10", "仍建议补通信模块消融、FDI 分类准确率、严重故障回退、高曲率相位角色、运行时开销、与 TF13 同场对齐统计。"),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    headers = ["状态", "缺口编号", "当前结论"]
    widths = [1400, 1700, 6260]
    for cell, header in zip(table.rows[0].cells, headers):
        set_cell_text(cell, header, bold=True, color=(31, 78, 121), fill="EAF2F8")
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            fill = "F6FBF7" if row[0] == "已部分补齐" else "FFF8E8"
            set_cell_text(cells[idx], value, fill=fill)
    set_table_width(table, widths)


def add_figure(doc, file_name, caption):
    image_path = GAP_FIG_DIR / file_name
    if not image_path.exists():
        raise FileNotFoundError(image_path)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(image_path), width=Inches(6.15))

    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.space_after = Pt(8)
    c.paragraph_format.keep_together = True
    run = c.add_run(caption)
    set_run_font(run, size=9, name="宋体", color=(80, 80, 80))


def main():
    if not SRC_DOCX.exists():
        raise FileNotFoundError(SRC_DOCX)
    doc = Document(str(SRC_DOCX))

    replaced = []
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text in FORMULA_REPLACEMENTS:
            image = EQ_DIR / FORMULA_REPLACEMENTS[text]
            if not image.exists():
                raise FileNotFoundError(image)
            insert_equation_image(paragraph, image)
            replaced.append(text)

    missing = set(FORMULA_REPLACEMENTS) - set(replaced)
    if missing:
        print("WARNING: formulas not replaced:")
        for item in sorted(missing):
            print(" -", item)

    add_heading(doc, "17. Final gap-closure 图包评估与入稿更新", level=1)
    add_body_paragraph(
        doc,
        "本节根据 tf14_final_gap_closure_20260509 更新。该图包新增 6 张补缺口图，分别覆盖 Koopman 预测验证、positive=worse 消融热力图、模态证书裕度统计，以及力变化率、jerk、角点载荷时间线。总体判断是：这些图可以作为当前中文稿的阶段性证据和补充材料候选，用来把方法模块与实验支撑对应起来；但它们不能替代最终 n≥20 多随机种子统计。",
    )
    add_body_paragraph(
        doc,
        "写作时需要保持边界清楚：E1 来自缓存离线筛选数据，不等价于完整 DNN 多种子重训练；E2、E3、E7 当前主要用于说明指标定义、消融趋势和力学诊断链路已经闭合。正文中可以用“阶段性诊断结果”“补充材料候选”“最终统计仍待 n≥20 扩展”等表述，避免把单次运行图说成最终显著性结论。",
    )

    add_heading(doc, "17.1 可用性判断", level=2)
    add_gap_table(doc)

    add_heading(doc, "17.2 新增图件", level=2)
    for file_name, caption in GAP_FIGURES:
        add_figure(doc, file_name, caption)

    add_heading(doc, "17.3 与原缺口清单的关系", level=2)
    add_status_table(doc)
    add_body_paragraph(
        doc,
        "据此，中文稿中的实验论证可以从“还缺少若干关键图”更新为“E1/E2/E3/E7 已形成阶段性证据闭环，剩余工作主要是多种子统计与若干专项模块验证”。这一改法对顶刊审稿更稳妥：既展示 TF14 已经把 Koopman 建模、通信鲁棒消融、稳定性证书和力学安全诊断串起来，又主动承认当前图包仍属于 gap-closure 版本，而不是最终统计闭环。",
    )

    doc.save(str(OUT_DOCX))
    print(OUT_DOCX)
    print("formula_replaced", len(replaced))
    print("inline_shapes", len(doc.inline_shapes))
    print("tables", len(doc.tables))


if __name__ == "__main__":
    main()
