from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "paper_dcn_tf12_draft"
OUT_DOCX = OUT_DIR / "tf14_method_update_zh_2026-05-08.docx"

SKILL_SCRIPTS = Path(
    r"C:\Users\lj\.codex\plugins\cache\openai-primary-runtime\documents\26.430.10722"
    r"\skills\documents\scripts"
)
sys.path.append(str(SKILL_SCRIPTS))
from table_geometry import apply_table_geometry, column_widths_from_weights  # noqa: E402


ACCENT = RGBColor(23, 78, 111)
MUTED = RGBColor(93, 103, 112)
LIGHT_FILL = "EAF2F6"
NOTE_FILL = "F6F8FA"
TABLE_HEADER_FILL = "D9EAF2"


def set_east_asia_font(run, east_asia: str = "Microsoft YaHei", latin: str = "Arial"):
    run.font.name = latin
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)


def set_style_font(style, size_pt=None, bold=None, color=None):
    font = style.font
    font.name = "Arial"
    if size_pt is not None:
        font.size = Pt(size_pt)
    if bold is not None:
        font.bold = bold
    if color is not None:
        font.color.rgb = color
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def shade_cell(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text: str, bold=False, color=None, size=9.5, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    set_east_asia_font(run)
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_table(doc, headers, rows, weights, font_size=9):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for idx, h in enumerate(headers):
        shade_cell(hdr.cells[idx], TABLE_HEADER_FILL)
        set_cell_text(hdr.cells[idx], h, bold=True, color=ACCENT, size=font_size)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            set_cell_text(cells[idx], str(value), size=font_size)
    widths = column_widths_from_weights(weights, total_width_dxa=9360)
    apply_table_geometry(table, widths, table_width_dxa=9360, indent_dxa=0)
    doc.add_paragraph()
    return table


def add_para(doc, text: str = "", style: str | None = None, bold=False, color=None):
    p = doc.add_paragraph(style=style)
    if text:
        r = p.add_run(text)
        set_east_asia_font(r)
        r.bold = bold
        if color is not None:
            r.font.color.rgb = color
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(5)
        r = p.add_run(item)
        set_east_asia_font(r)
        r.font.size = Pt(10.5)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(5)
        r = p.add_run(item)
        set_east_asia_font(r)
        r.font.size = Pt(10.5)


def add_equation(doc, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    set_east_asia_font(r, east_asia="Microsoft YaHei", latin="Consolas")
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor(35, 43, 50)


def add_note_box(doc, title: str, body: str):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    shade_cell(cell, NOTE_FILL)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    set_east_asia_font(r)
    r.bold = True
    r.font.color.rgb = ACCENT
    r.font.size = Pt(10.5)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run(body)
    set_east_asia_font(r2)
    r2.font.size = Pt(10)
    apply_table_geometry(table, [9360], table_width_dxa=9360, indent_dxa=0)
    doc.add_paragraph()


def add_figure(doc, rel_path: str, caption: str, width_in=6.25):
    path = ROOT / rel_path
    if not path.exists():
        add_note_box(doc, "图位暂缺", f"{caption}。当前未找到文件：{path}")
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width_in))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(8)
    r = cap.add_run(caption)
    set_east_asia_font(r)
    r.italic = True
    r.font.size = Pt(9)
    r.font.color.rgb = MUTED


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r1 = paragraph.add_run("第 ")
    set_east_asia_font(r1)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    r2 = paragraph.add_run()
    r2._r.append(fld_begin)
    r2._r.append(instr)
    r2._r.append(fld_end)
    r3 = paragraph.add_run(" 页")
    set_east_asia_font(r3)


def setup_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    styles = doc.styles
    set_style_font(styles["Normal"], 10.5)
    styles["Normal"].paragraph_format.line_spacing = 1.12
    styles["Normal"].paragraph_format.space_after = Pt(5)
    set_style_font(styles["Title"], 20, True, ACCENT)
    set_style_font(styles["Subtitle"], 11, False, MUTED)
    set_style_font(styles["Heading 1"], 15, True, ACCENT)
    set_style_font(styles["Heading 2"], 12.5, True, RGBColor(30, 63, 82))
    set_style_font(styles["Heading 3"], 11, True, RGBColor(30, 63, 82))
    for style_name in ("List Bullet", "List Number"):
        set_style_font(styles[style_name], 10.5)
        styles[style_name].paragraph_format.left_indent = Inches(0.28)
        styles[style_name].paragraph_format.first_line_indent = Inches(-0.12)
        styles[style_name].paragraph_format.space_after = Pt(5)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hr = header.add_run("TF14 中文方法更新稿 | Network-Resilient Cooperative Transport Control")
    set_east_asia_font(hr)
    hr.font.size = Pt(8.5)
    hr.font.color.rgb = MUTED

    footer = section.footer.paragraphs[0]
    add_page_number(footer)
    for run in footer.runs:
        run.font.size = Pt(8.5)
        run.font.color.rgb = MUTED

    return doc


def build_docx() -> None:
    doc = setup_document()

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = title.add_run("TF14 方法更新中文稿")
    set_east_asia_font(tr)
    tr.bold = True
    tr.font.color.rgb = ACCENT
    tr.font.size = Pt(22)

    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = subtitle.add_run(
        "面向 Digital Communications and Networks 投稿的主方法修订版："
        "在线 FDI、切换容错、阶段角色调度与稳定性证书"
    )
    set_east_asia_font(sr)

    add_table(
        doc,
        ["项目", "内容"],
        [
            ["依据文件", "tf14_pre.ipynb、tf14_runtime.py、control_files/tf14/*、Stage4-Stage6 实验摘要"],
            ["稿件定位", "新的中文方法版本，不覆盖旧中文稿；用于替换和强化论文主方法与实验设计逻辑"],
            ["核心变化", "从 TF13 的已知故障补偿，升级为未知故障下的在线诊断、辨识、切换容错与证书化验证闭环"],
            ["生成日期", "2026-05-08"],
        ],
        [1.2, 5.8],
        font_size=9.5,
    )

    add_note_box(
        doc,
        "写作定位",
        "TF14 不应被写成“又调了一版参数”。它的主要价值是把原先依赖预设故障信息的容错控制，"
        "改成在线可检测、可辨识、可切换、可解释并可用 Lyapunov 型证书监控的网络韧性协同运输控制框架。"
        "因此论文贡献、方法公式、消融实验和图表都必须围绕这条闭环展开。",
    )

    doc.add_heading("1. TF14 相对旧稿的定位", level=1)
    add_para(
        doc,
        "本文系统仍以四车协同运输刚性/半刚性载荷为对象，基础控制框架保留双线性 Koopman 升维预测和 MPC。"
        "最新 TF14 系列的实质更新在于：控制器不再假设故障车辆和故障幅值已知，而是在运行中通过模型残差和执行器效率估计识别故障，"
        "再根据故障严重度、通信质量和安全证书切换控制模式。这样，方法贡献从“预测模型更准”扩展为“网络扰动和未知执行器退化下的自治容错协同运输”。",
    )
    add_bullets(
        doc,
        [
            "相对 baseline/AKE 思路：baseline 重点是自适应 Koopman 表示与鲁棒控制，没有把多车协同运输中的载荷内力、通信质量、未知执行器退化和在线容错切换放在同一闭环中处理。",
            "相对 TF13：TF13 已有稳定投影双线性 Koopman、通信补偿和故障重分配，但故障更接近“已知注入后补偿”；TF14 增加在线 FDI、效率辨识、切换容错和证书记录，论文叙事必须从已知故障补偿转向未知故障自治韧性。",
            "相对普通 MPC 或多车一致性控制：TF14 的 MPC 不是孤立跟踪器，而是由 Koopman 预测、通信质量一致性、时延补偿、阶段角色调度、故障重构和安全证书共同约束的分层控制器。",
        ]
    )

    add_table(
        doc,
        ["维度", "baseline / TF13 的不足", "TF14 的新增机制", "内在机理"],
        [
            [
                "模型",
                "baseline 主要强调可学习动力学；TF13 已引入双线性结构但仍以主控制性能为核心",
                "稳定投影双线性 Koopman + 离线 ridge 拟合 + 可选在线 bilinear ridge 残差修正",
                "双线性项显式表达控制输入与升维状态的乘性交互，更贴近轮胎侧偏、曲率和输入耦合；稳定投影抑制长时域预测发散。",
            ],
            [
                "故障信息",
                "TF13 容错依赖配置中已知故障开始、故障车辆和缩放比例",
                "FDI 残差、CUSUM、转向/加速度效率估计与置信度融合",
                "先判断实际执行能力是否下降，再触发补偿，避免把通信抖动或短时模型误差误当成真实执行器故障。",
            ],
            [
                "容错控制",
                "旧版主要做故障后补偿或力矩重分配，缺少模式化安全边界",
                "nominal、reconfigured FTC、safe degraded consensus 三模式切换",
                "效率逆补偿恢复故障车辆的等效指令，支持车辆吸收缺失力/力矩；严重退化或低通信质量下进入保守模式，降低过激控制风险。",
            ],
            [
                "协同角色",
                "车辆被近似同质处理，缺少弯道阶段和前后/内外侧角色差异",
                "阶段-角色调度器按 straight、turn entry/core/exit/transition 生成有界小修正",
                "在弯道内外侧和前后桥承担不同载荷/航向误差时，小幅修正可降低载荷偏航和单车误差集中。",
            ],
            [
                "可证明性",
                "实验曲线能说明效果，但安全性和稳定性论证弱",
                "切换 Lyapunov 型证书记录 V、收缩裕度和 certificate_ok",
                "把通信质量、辨识误差和切换模式写入 Lyapunov 上界，使仿真实验能直接支撑稳定性定理的假设和结论。",
            ],
        ],
        [1.1, 2.0, 2.0, 2.9],
        font_size=8.2,
    )

    doc.add_heading("2. 系统建模：从大地坐标到路径误差坐标", level=1)
    add_para(
        doc,
        "方法部分建议先从单车动力学和坐标系写起，再进入多车载荷约束。车辆在大地坐标系中的质心状态可写为 "
        "q_i = [X_i, Y_i, psi_i, v_x,i, v_y,i, r_i]^T，控制输入为 u_i = [delta_i, a_x,i]^T。"
        "最基础的全局运动学关系为：",
    )
    add_equation(
        doc,
        "dot X_i = v_x,i cos psi_i - v_y,i sin psi_i,  "
        "dot Y_i = v_x,i sin psi_i + v_y,i cos psi_i,  dot psi_i = r_i."
    )
    add_para(
        doc,
        "为便于路径跟踪和 Koopman-MPC 建模，将全局状态投影到参考路径的 Frenet 坐标，得到 "
        "x_i = [s_i, e_y,i, e_psi,i, v_x,i, v_y,i, r_i]^T。若参考路径曲率为 kappa(s)，"
        "则路径方向误差系统为：",
    )
    add_equation(
        doc,
        "dot s_i = (v_x,i cos e_psi,i - v_y,i sin e_psi,i)/(1 - kappa(s_i)e_y,i)"
    )
    add_equation(
        doc,
        "dot e_y,i = v_x,i sin e_psi,i + v_y,i cos e_psi,i,  "
        "dot e_psi,i = r_i - kappa(s_i) dot s_i"
    )
    add_para(
        doc,
        "横向速度、横摆角速度和纵向速度采用线性轮胎近似下的动态自行车模型，其中 m、I_z、l_f、l_r、C_f、C_r 分别为车辆质量、横摆转动惯量、前/后轴距和前/后轮侧偏刚度：",
    )
    add_equation(
        doc,
        "dot v_y,i = -(2C_f+2C_r)/(m v_x,i) v_y,i + "
        "(-v_x,i -(2C_f l_f-2C_r l_r)/(m v_x,i)) r_i + (2C_f/m) delta_i"
    )
    add_equation(
        doc,
        "dot r_i = -(2C_f l_f-2C_r l_r)/(I_z v_x,i) v_y,i - "
        "(2C_f l_f^2+2C_r l_r^2)/(I_z v_x,i) r_i + (2C_f l_f/I_z) delta_i"
    )
    add_equation(doc, "dot v_x,i = a_x,i + r_i v_y,i,  v_x,i >= 0.5 m/s.")
    add_para(
        doc,
        "多车协同运输层把每辆车的连接点和载荷参考中心绑定。若载荷中心位姿为 p_L, psi_L，车辆 i 的期望连接偏置为 ell_i，"
        "则连接误差 e_c,i = p_i - p_L - R(psi_L) ell_i。载荷约束项同时进入 MPC 代价和安全约束，用来抑制车间距离漂移、"
        "连接拉伸以及载荷内力峰值。后续图表应单独给出载荷合力、纵向/横向分力和单车连接误差，否则“协同运输”证据不足。",
    )

    doc.add_heading("3. 离线数据与 Koopman 升维学习", level=1)
    add_para(
        doc,
        "离线阶段使用上述车辆-路径误差模型生成覆盖不同速度、曲率、横向误差、航向误差和输入幅值的数据。故障与通信扰动不建议直接混入主 Koopman 训练目标，"
        "而应作为在线鲁棒性层处理：这样可以让 Koopman 模型学习正常可控动力学，FDI/FTC 层专门处理异常执行能力和网络质量变化。",
    )
    add_para(
        doc,
        "TF14 的升维状态可写为 z_i = Phi_theta(x_i) = [x_i, phi_theta(x_i)]，其中 phi_theta 为深度编码器。"
        "最新 tf14_pre 配置中，编码器隐藏宽度为 128，深度为 4，输出维度为 24；训练轮数 180，batch size 4096，weight decay 为 8e-5。"
        "离线线性/双线性矩阵使用 ridge_lambda = 5e-5，并通过 linear_fit_blend = 0.60 与 previous_model_blend = 0.20 稳定继承已有模型。",
    )
    add_equation(
        doc,
        "z_{k+1} = A z_k + B u_k + sum_{j=1}^{2} u_{k,j} N_j z_k + epsilon_k."
    )
    add_para(
        doc,
        "相比单纯线性 Koopman，双线性结构把输入和状态之间的乘性交互显式保留下来。车辆控制中，转角对横摆响应的影响会随速度、侧偏状态和曲率变化而变化；"
        "如果只使用 z_{k+1}=Az_k+Bu_k，控制影响被固定为常数矩阵，容易在高曲率和故障工况下出现系统性偏差。"
    )
    add_para(
        doc,
        "在线自适应接口仍保留 bilinear ridge 残差更新，典型配置为窗口 18、每 6 步更新、forget_factor=0.97、ridge_lambda=2e-4、"
        "dA/dB 范数上限 0.06、blend=0.04、投影半径 0.998。需要注意：实时主 preset 中为了保证控制预算，默认启用实时模式并关闭在线更新；"
        "若论文中讨论在线自适应收益，应在非实时或单独消融 preset 中开启，而不能把实时主实验和在线更新收益混写。",
    )

    doc.add_heading("4. TF14 在线 FDI 与执行器效率辨识", level=1)
    add_para(
        doc,
        "TF14 的第一个核心新增模块是 OnlineFDIMonitorTF14。它对每辆车做一阶预测，比较预测状态和实际状态，得到归一化残差；同时估计转向和纵向加速度执行效率。"
        "残差用于发现异常，效率估计用于区分故障类型，通信质量用于降低误报风险。",
    )
    add_equation(
        doc,
        "hat x_{i,k|k-1}=f_i(x_{i,k-1},u^{cmd}_{i,k-1}),  "
        "r_{i,k}=x_{i,k}-hat x_{i,k|k-1}"
    )
    add_equation(
        doc,
        "rho_{i,k}=sqrt(sum_l (r_{i,k,l}/eta_l)^2),  "
        "EWMA_k=beta EWMA_{k-1}+(1-beta)rho_{i,k}"
    )
    add_equation(
        doc,
        "u^{act}_{i,k}=Gamma_{i,k}u^{cmd}_{i,k},  "
        "Gamma_{i,k}=diag(gamma_delta,i,k, gamma_a,i,k)."
    )
    add_para(
        doc,
        "当可获得实际执行量时，效率可由 u_act/u_cmd 直接估计；否则用 v_x 加速度和横摆角速度响应反推。"
        "当前阈值设置包括 residual_threshold=0.85、CUSUM 阈值 8.0、纵向效率故障阈值 0.78、转向效率故障阈值 0.80、严重退化阈值 0.35。"
        "检测需要 hold_steps=5，释放需要 release_hold_steps=10，以避免抖振。",
    )
    add_bullets(
        doc,
        [
            "residual_only_alert：残差升高但执行效率尚未确认下降，通常用于模型偏差、通信抖动或短时异常的预警。",
            "accel_degraded / steer_degraded：纵向或转向执行效率低于阈值，触发对应方向的补偿和重分配。",
            "dual_degraded / severe_dual_loss：两类执行器同时退化或严重下降，后续控制器会进入重构容错或安全降级模式。",
            "通信质量低于软阈值时，残差阈值会保守调整，避免把网络干扰与物理执行器故障混淆。",
        ]
    )
    add_figure(
        doc,
        "results/tf14_diagnostics/tf14_fdi_identification.png",
        "图 1  在线 FDI 与执行器效率辨识诊断图：用于支撑“未知故障可被检测并分类”的贡献。",
    )

    doc.add_heading("5. 切换容错控制：从诊断到重构", level=1)
    add_para(
        doc,
        "ReconfigurableFTCControllerTF14 将 MPC 输出 u_MPC 作为基准指令，再根据 FDI 结果、通信质量和 dwell 约束选择三种控制模式。"
        "该模块是 TF14 相对 TF13 最需要写清楚的地方：它不只是给故障车加一个补偿项，而是把诊断置信度、执行器效率、支持车辆重分配和安全降级统一成切换系统。",
    )
    add_numbered(
        doc,
        [
            "nominal_koopman_mpc：未发现可信故障时直接执行 Koopman-MPC 输出，即 u_i = u_i^MPC。",
            "reconfigured_ftc_mpc：故障置信度超过 activation_confidence=0.12 后，对故障车做效率逆补偿，并把无法完成的等效力/力矩缺口分配给支持车辆。",
            "safe_degraded_consensus：当双执行器严重退化或全局通信质量低于约 0.18 时，进入安全降级一致性控制，对全队转角和加速度做保守缩放。",
        ]
    )
    add_equation(
        doc,
        "u_i^R = Pi_U( hat Gamma_i^{-1} u_i^MPC + Delta u_i^{redist} ),  "
        "sum_i Delta u_i^{redist} approx missing team effort."
    )
    add_equation(
        doc,
        "u_i^S = Pi_U( diag(alpha_delta, alpha_a) u_i^MPC + Delta u_i^{safe} )."
    )
    add_para(
        doc,
        "这里的 Pi_U 是输入约束投影，防止逆补偿把指令推到不可执行区域。switch_dwell_steps=4 用来限制频繁切换。"
        "该设计的物理机理是：故障车辆的实际输出下降会导致队形和载荷力矩不平衡，逆补偿先尝试恢复故障车局部能力；若恢复不足，支持车辆承担剩余缺口；"
        "若通信或执行器状态已经无法可信闭环，则降级为低激进度控制以保住安全边界。",
    )
    add_figure(
        doc,
        "results/tf14_diagnostics/tf14_switch_certificate.png",
        "图 2  切换模式与 Lyapunov 型证书：用于展示故障后从 nominal 到 reconfigured FTC 的模式转移及证书裕度。",
    )

    doc.add_heading("6. 阶段-角色调度器", level=1)
    add_para(
        doc,
        "PhaseRoleSchedulerTF14 是 TF14 后续阶段中新增的协同运输细化模块。它并不替代 MPC，而是在 FTC 输出之后加入有界小修正：",
    )
    add_equation(
        doc,
        "u_i^{out}=Pi_U(u_i^{FTC} + [Delta delta_i^{phase}, Delta a_i^{phase}]^T)."
    )
    add_para(
        doc,
        "调度器先按参考曲率及其变化把路径划分为 straight、turn_entry、turn_core、turn_exit 和 turn_transition，"
        "再根据车辆角色 FL、FR、RL、RR 以及前/后、内/外侧关系修正转角和加速度。"
        "当通信质量下降时，调度器会削弱转角修正并收缩加速度 trim；当某辆车被 FDI 判定为故障车时，故障车修正缩小，支持车辆承担更多补偿。",
    )
    add_para(
        doc,
        "这一模块的创新点不能只写“加入规则控制”。它的贡献在于把协同运输中的路径阶段和车辆力学角色显式编码，让不同车辆在弯道入口、弯中和出弯承担不同的横摆/纵向误差修正任务。"
        "如果图表能证明它降低载荷偏航、连接拉伸或单车峰值误差，即可形成比普通同质多车 MPC 更强的协同运输叙事。",
    )
    add_figure(
        doc,
        "tf14_stage5_phase_role_scenarios_20260507/figures/dlc_mixed_fault_noise_phase_role_trim.png",
        "图 3  混合故障与通信噪声下的阶段-角色 trim：用于说明调度器确实在不同路径阶段输出有界修正。",
    )

    doc.add_heading("7. MPC、安全约束与实时预算", level=1)
    add_para(
        doc,
        "MPC 使用 Koopman 预测模型在有限时域内优化路径跟踪、输入平滑、队形连接和载荷安全项。当前主配置 horizon=12，max_sqp_iters=1，"
        "转角上限约 0.52 rad，加速度上限约 4.0 m/s^2，并使用曲率相关速度规划、终端减速、PPC/dynamic PPC、进度监督和紧急保护。"
    )
    add_para(
        doc,
        "通信韧性层包括通信质量一致性、时延补偿、通信约束收紧和 degraded fallback。可在论文中把通信质量 q_k 写入约束收紧系数，"
        "例如 e_max(q_k)=e_max^0 - c_q(1-q_k)，或把队形/输入平滑权重写成 q_k 的函数。这样贡献就能与 Digital Communications and Networks 的主题更直接对齐。",
    )
    add_para(
        doc,
        "实时预算是 TF14 的一个方法取舍。默认实时 preset 使用 leader-first、mpc_decimation_steps=5、每步至少求解 2 辆车、可在预算不足时跳过部分 MPC，"
        "并将控制预算设为约 0.018 s。Stage6 的 error-match preset 则改为每步求解全部 4 辆车，mpc_decimation_steps=1、mpc_skip_on_budget=False，"
        "牺牲单步时间换取与 TF13 更接近的横向误差。论文里应把这两个 preset 分开：主方法强调实时韧性，补充实验展示精度-计算时间前沿。",
    )

    doc.add_heading("8. Lyapunov 型稳定性证书", level=1)
    add_para(
        doc,
        "TF14 的稳定性证明建议写成切换系统的输入到状态实用稳定性。令团队误差 e_k = x_k - x_ref,k，切换模式 sigma_k 属于 "
        "{N, R, S}，分别对应 nominal、reconfigured FTC 和 safe degraded consensus。闭环误差可抽象为：",
    )
    add_equation(
        doc,
        "e_{k+1}=A_{sigma_k}e_k + B_{sigma_k} tilde Gamma_k u_k + d_k,"
    )
    add_para(
        doc,
        "其中 tilde Gamma_k 表示真实执行器效率与估计效率之间的残余误差，d_k 包含模型误差、通信扰动和时延补偿残差。"
        "对每个模式构造二次函数 V_sigma(e)=e^T P_sigma e，并假设诊断/切换满足平均驻留时间，扰动和辨识误差有界，则可得到：",
    )
    add_equation(
        doc,
        "V_{sigma_{k+1}}(e_{k+1}) <= (1 - lambda_sigma Delta t) V_{sigma_k}(e_k) "
        "+ c_d ||d_k||^2 + c_g ||Gamma_k - hat Gamma_k||^2."
    )
    add_para(
        doc,
        "这说明：只要 FDI 在有限时间内给出有界误差的效率估计，FTC 切换不发生无限快抖振，且通信扰动有界，团队跟踪误差最终进入一个由扰动上界和辨识误差上界决定的小邻域。"
        "代码中的 SwitchedFTCCertificateTF14 正是在每一步记录 V、predicted_upper、contraction_margin 和 certificate_ok；它是证明假设与仿真实验之间的桥梁，而不是单独替代理论证明。",
    )

    doc.add_heading("9. 现有 TF14 证据可以怎样写入论文", level=1)
    add_table(
        doc,
        ["阶段", "当前证据", "可写入论文的结论", "仍需补强"],
        [
            [
                "Stage3",
                "理论-实现映射覆盖 23 项，覆盖率 100%",
                "证明链条中的 FDI、切换、证书、通信项在代码中都有对应实现",
                "该结果只能说明实现覆盖，不等于最终理论证明。",
            ],
            [
                "Stage4",
                "推荐默认候选达到平均 1 step 检测、1 step 切换、0 误报、完整路径率 1.0",
                "可支撑在线 FDI/FTC 的快速响应能力",
                "需要更复杂故障：慢变、间歇、偏置、双车同时故障。",
            ],
            [
                "Stage5",
                "DLC 混合故障/通信噪声中，baseline RMSE_y=0.0453，TF14 main=0.0409，phase-role=0.0404；RMSE_s 约 0.4711 -> 0.4677",
                "可说明在混合扰动下横向误差和纵向误差均有改善，阶段角色调度有小幅附加收益",
                "单次实验说服力不足，需要多 seed 统计和载荷力/连接误差图。",
            ],
            [
                "Stage6",
                "全车每步求解 preset 的平均横向 RMSE 相对 TF13 目标差距 +1.59%，纵向 RMSE 低约 47.80%",
                "可作为精度-计算时间前沿实验，说明 TF14 可切换到高精度模式",
                "需要同步展示 step_mean 上升到约 0.071-0.075 s 的代价。",
            ],
        ],
        [0.9, 2.4, 2.4, 2.1],
        font_size=8.2,
    )
    add_figure(
        doc,
        "tf14_stage4_fdi_fast_response_20260507/figures/tf14_stage4_fdi_tuning_summary.png",
        "图 4  FDI/FTC 快速响应调参摘要：建议作为补充材料或方法可靠性图。",
    )
    add_figure(
        doc,
        "tf14_stage5_phase_role_scenarios_20260507/figures/dlc_mixed_fault_noise_system_lat_long_error_compare.png",
        "图 5  DLC 混合故障通信噪声下系统级横向/纵向误差对比：可作为 TF14 主实验候选图。",
    )
    add_figure(
        doc,
        "tf14_stage5_phase_role_scenarios_20260507/figures/tf14_stage5_metrics_summary.png",
        "图 6  Stage5 多工况指标汇总：可用于展示 baseline、TF14 main 和 TF14 phase-role 的整体趋势。",
    )

    doc.add_heading("10. 新增实验设计一：未知执行器故障诊断-切换-证书闭环", level=1)
    add_para(
        doc,
        "该实验专门验证 TF14 最核心的新增贡献：未知故障发生后，系统是否能在线检测、辨识、切换、重分配，并保持 Lyapunov 型证书不失效。"
        "它应作为主实验或主消融，而不是补充材料。",
    )
    add_table(
        doc,
        ["要素", "设计"],
        [
            [
                "对比方法",
                "baseline/AKE-MPC；TF13 known-fault compensation；TF14 no-FDI/no-switch；TF14 FDI-only；TF14 full；TF14 full + fast-response FDI。",
            ],
            [
                "故障类型",
                "单车纵向效率下降、单车转向效率下降、双执行器同时下降、严重双执行器退化、慢变 ramp 故障、间歇 dropout、多车同时故障。",
            ],
            [
                "通信条件",
                "clean、medium loss/delay、高丢包高时延、故障与通信噪声同时发生。故障开始时间随机化，不只固定在 s=24 m。",
            ],
            [
                "核心指标",
                "检测延迟、切换延迟、误报率、故障类型识别准确率、certificate_ok ratio、最小 contraction margin、RMSE_y/RMSE_s、max_y/max_s、载荷合力和内力峰值。",
            ],
            [
                "必须出图",
                "FDI 残差/效率时间线；模式切换时间线；故障车与支持车控制重分配；Lyapunov V 与 margin；系统与单车误差；载荷纵向/横向力和合力。",
            ],
            [
                "预期结论",
                "TF14 full 应在未知故障下显著降低故障后误差峰值和载荷力峰值，同时保持低误报；no-FDI/no-switch 应暴露出无法自治触发容错的问题。",
            ],
        ],
        [1.25, 5.75],
        font_size=9,
    )
    add_note_box(
        doc,
        "这组实验要证明什么",
        "审稿人不会只看最终 RMSE，而会问：故障是不是未知？诊断是不是比控制收益更早发生？切换后谁承担了缺失控制量？"
        "证书是否在故障后仍成立？因此图必须按“故障发生 -> FDI 置信度升高 -> 模式切换 -> 重分配 -> 误差/力恢复 -> 证书保持”这一时间链组织。",
    )

    doc.add_heading("11. 新增实验设计二：阶段-角色调度与实时预算/精度前沿", level=1)
    add_para(
        doc,
        "第二套实验用于把 Stage5 和 Stage6 的结果组织成更有说服力的论文实验：一方面证明阶段-角色调度不是装饰模块，另一方面解释实时 preset 与全车求解 preset 的取舍。",
    )
    add_table(
        doc,
        ["要素", "设计"],
        [
            [
                "路径/场景",
                "DLC、hairpin、S-curve、窄通道避障；每个场景分别设置 clean、high communication degradation、single fault、mixed fault/noise。",
            ],
            [
                "对比方法",
                "TF14 main without phase-role；TF14 phase-role；TF14 phase-role with trim disabled per role；TF14 realtime-decimated；TF14 all-vehicle-solve error-match preset。",
            ],
            [
                "核心指标",
                "RMSE_y/RMSE_s、单车峰值误差、载荷 yaw/横摆相平面面积、连接误差、载荷纵向/横向力峰值、mean step time、solver overrun ratio、每步求解车辆数。",
            ],
            [
                "必须出图",
                "路径阶段标签图；四车 phase-role trim；yaw phase plane；误差-计算时间 Pareto 图；单车横向/纵向误差热图；载荷力/力矩对比图。",
            ],
            [
                "预期结论",
                "phase-role 在高曲率和混合扰动下应降低内外侧车辆误差差异或载荷力峰值；all-vehicle solve 可逼近 TF13 横向精度但计算时间更高，realtime preset 更适合作为主实时控制。",
            ],
        ],
        [1.25, 5.75],
        font_size=9,
    )
    add_note_box(
        doc,
        "这组实验要证明什么",
        "阶段-角色调度的价值不一定体现在平均 RMSE 大幅下降，而可能体现在单车峰值误差、载荷内力、横摆相平面面积和高曲率段稳定性。"
        "如果只放系统平均误差，审稿人会认为该模块收益很弱；必须用载荷与角色相关图证明其物理意义。",
    )

    doc.add_heading("12. 需要同步补齐的图表清单", level=1)
    add_bullets(
        doc,
        [
            "离线数据覆盖图：s、e_y、e_psi、v_x、delta、a_x、曲率的分布，用来证明 Koopman 训练数据覆盖控制域。",
            "Koopman 预测误差图：one-step 与 multi-step 误差，至少比较 linear Koopman、bilinear Koopman、bilinear + stable projection。",
            "FDI 时间链图：故障真实发生时间、残差、CUSUM、gamma_delta/gamma_a、检测状态、故障类型、置信度。",
            "FTC 重分配图：故障车辆实际执行能力下降后，各车辆 delta/a_x 修正量和 redistributed_total_delta/ax。",
            "证书图：V_sigma、predicted_upper、contraction_margin、certificate_ok，并标出模式切换时刻。",
            "系统级轨迹和误差：payload/team center 轨迹、RMSE_y/RMSE_s、max_y/max_s，与 baseline 和 TF13/TF14 消融对比。",
            "单车级误差：四辆车的横向、纵向、航向误差，突出故障车和支持车差异。",
            "载荷力学图：总合力、纵向分力、横向分力、偏航力矩、连接误差和内力峰值。",
            "阶段-角色图：路径阶段标签、各车辆 trim、内外侧/前后角色在弯道中的控制分担。",
            "实时性图：step time 分布、solver time、overrun ratio、求解车辆数，与误差指标组成 Pareto 图。",
        ]
    )

    doc.add_heading("13. 建议改写后的四条贡献", level=1)
    add_table(
        doc,
        ["编号", "贡献表述"],
        [
            [
                "贡献 1",
                "提出面向多车协同运输的稳定投影双线性 Koopman-MPC 框架，用升维双线性模型描述车辆输入-状态耦合，并在路径误差、队形连接和载荷安全约束下实现预测控制。",
            ],
            [
                "贡献 2",
                "提出通信感知的在线 FDI 与执行器效率辨识机制，通过模型残差、CUSUM、转向/纵向效率估计和通信质量调节，实现未知故障的低延迟检测与类型识别。",
            ],
            [
                "贡献 3",
                "提出切换式容错 MPC，包括效率逆补偿、支持车辆重分配、安全降级一致性和 dwell-time 防抖，并给出模式依赖 Lyapunov 型稳定性证明与在线证书记录。",
            ],
            [
                "贡献 4",
                "提出阶段-角色协同调度和实时预算机制，在弯道阶段、前后/内外侧车辆角色、故障车/支持车关系和计算预算之间进行有界调节，形成精度-安全-实时性的可解释折中。",
            ],
        ],
        [1.0, 6.0],
        font_size=9,
    )

    add_note_box(
        doc,
        "投稿前的硬性提醒",
        "当前 TF14 已经具备更清楚的方法创新链，但顶刊/一区审稿会要求每个模块都有独立消融和物理量支撑。"
        "尤其是 FDI/FTC、phase-role、certificate 和 real-time budget 不能只用平均 RMSE 证明；必须同时给出时间链、力学量、稳定性证书和计算代价。",
    )

    doc.save(OUT_DOCX)


if __name__ == "__main__":
    build_docx()
    print(OUT_DOCX)
