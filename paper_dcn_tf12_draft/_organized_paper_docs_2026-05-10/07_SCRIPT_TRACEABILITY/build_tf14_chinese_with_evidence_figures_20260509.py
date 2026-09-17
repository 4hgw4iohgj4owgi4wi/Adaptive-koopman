from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
BASE_DOCX = ROOT / "paper_dcn_tf12_draft" / "tf14_method_update_zh_2026-05-08.docx"
OUT_DOCX = ROOT / "paper_dcn_tf12_draft" / "tf14_method_update_zh_with_evidence_figures_2026-05-09.docx"
FIG_DIR = ROOT / "tf14_paper_figures_20260508" / "figures"

SKILL_SCRIPTS = Path(
    r"C:\Users\lj\.codex\plugins\cache\openai-primary-runtime\documents\26.430.10722"
    r"\skills\documents\scripts"
)
sys.path.append(str(SKILL_SCRIPTS))
from table_geometry import apply_table_geometry, column_widths_from_weights  # noqa: E402


ACCENT = RGBColor(20, 83, 112)
MUTED = RGBColor(91, 101, 110)
HEADER_FILL = "DDEFF6"
NOTE_FILL = "F5F7F9"
WARN_FILL = "FFF3CD"


def set_east_asia_font(run, east_asia: str = "Microsoft YaHei", latin: str = "Arial"):
    run.font.name = latin
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)


def content_width_dxa(doc: Document) -> int:
    sec = doc.sections[-1]
    return int(sec.page_width.twips - sec.left_margin.twips - sec.right_margin.twips)


def shade_cell(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text: str, *, bold=False, color=None, size=8.4, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(1.5)
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    set_east_asia_font(run)
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_table(doc: Document, headers, rows, weights, font_size=8.0, header_size=8.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for idx, h in enumerate(headers):
        shade_cell(hdr.cells[idx], HEADER_FILL)
        set_cell_text(hdr.cells[idx], h, bold=True, color=ACCENT, size=header_size)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            set_cell_text(cells[idx], str(value), size=font_size)
    total = content_width_dxa(doc)
    widths = column_widths_from_weights(weights, total_width_dxa=total)
    apply_table_geometry(table, widths, table_width_dxa=total, indent_dxa=0)
    doc.add_paragraph()
    return table


def add_para(doc: Document, text: str = "", style: str | None = None, bold=False, color=None):
    p = doc.add_paragraph(style=style)
    if text:
        r = p.add_run(text)
        set_east_asia_font(r)
        r.bold = bold
        if color is not None:
            r.font.color.rgb = color
    return p


def add_note_box(doc: Document, title: str, body: str, fill: str = NOTE_FILL):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    shade_cell(cell, fill)
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title)
    set_east_asia_font(r)
    r.bold = True
    r.font.size = Pt(10)
    r.font.color.rgb = ACCENT
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run(body)
    set_east_asia_font(r2)
    r2.font.size = Pt(9)
    total = content_width_dxa(doc)
    apply_table_geometry(table, [total], table_width_dxa=total, indent_dxa=0)
    doc.add_paragraph()


def add_figure(doc: Document, file_name: str, caption: str, width_in=6.6):
    path = FIG_DIR / file_name
    if not path.exists():
        add_note_box(doc, "图文件缺失", f"未找到：{path}", WARN_FILL)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width_in))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(8)
    r = cap.add_run(caption)
    set_east_asia_font(r)
    r.italic = True
    r.font.size = Pt(9)
    r.font.color.rgb = MUTED


def set_heading_style(doc: Document):
    for name in ("Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[name]
        style.font.name = "Arial"
        style.element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.color.rgb = ACCENT


def build_docx() -> None:
    doc = Document(str(BASE_DOCX))
    set_heading_style(doc)

    doc.add_page_break()
    add_para(doc, "14. 基于 2026-05-09 新图包的实验证据更新", style="Heading 1")
    add_para(
        doc,
        "本节根据 tf14_paper_figures_20260508/figures 中 2026-05-09 生成的图重新筛选。"
        "入文原则是：只放能够正向支撑控制效果、故障容错、安全约束或实时性的图；"
        "对仍属于草稿、占位或证据等级不足的图，只在表格中说明，不作为正证据使用。",
    )

    add_table(
        doc,
        ["图号", "是否放入中文稿", "可证明的结论", "使用时必须注意"],
        [
            [
                "F5 主轨迹对比",
                "放入",
                "在高通信退化和混合故障/噪声下，TF14 main 与 phase-role 能保持完整路径跟踪，轨迹没有发散。",
                "轨迹差异本身较小，不能单独证明性能提升；必须与 F6/S11 的误差图配合使用。",
            ],
            [
                "F6 跟踪误差统计",
                "放入",
                "TF14 在多个 DLC 场景下横向 RMSE 低于 baseline，代表性 e_y/e_s 曲线展示故障后恢复能力。",
                "Stage2 统计仍为 n=6；最终投稿需替换为 n>=20 的统计图。",
            ],
            [
                "F8 在线 FDI",
                "放入",
                "单车未知故障下，FDI 能在约 0.20 s 检测、0.26 s 触发 FTC 切换，且该组图中误报为 0。",
                "还缺故障类型混淆矩阵；目前更适合作为机制证据而非分类精度最终证据。",
            ],
            [
                "F9 FTC 切换与重分配",
                "放入",
                "故障后控制器从 nominal 切换到 reconfigured，车辆级 delta/a_x 指令保持在约束内，并出现重分配补偿。",
                "需要在最终稿中补充 severe fault/safe fallback 场景，否则 safe 模式只停留在设计层。",
            ],
            [
                "F10 载荷力与连接安全",
                "放入",
                "TF14 显著降低连接 stretch、连接 RMS/offset 和连接裕度占用率，能支撑协同运输安全性改善。",
                "TF14 的部分载荷力峰值更高，不能写成“受力全面降低”；应表述为“通过更主动力/力矩分配换取连接安全”。",
            ],
            [
                "S11 全场景误差曲线",
                "放入",
                "补充展示 clean、communication degradation、single fault、mixed fault/noise 四类场景下的 e_y/e_s 误差趋势。",
                "目前仍是代表性曲线，最终最好增加 mean±CI。",
            ],
            [
                "S15 控制饱和检查",
                "放入",
                "四车转角和加速度输入均未触碰设定边界，支撑 FTC 补偿没有通过不可执行输入获得效果。",
                "建议作为安全约束证据，放主文或补充材料均可。",
            ],
            [
                "S19 实时性对比",
                "放入",
                "TF14 main/phase-role 的平均步时接近 0.03 s，全路径成功率为 1；error-match/all-solve 精度更高但计算代价明显更大。",
                "这是实时性与精度取舍证据，不是单纯控制误差优势证据。",
            ],
        ],
        [0.9, 1.2, 3.0, 3.1],
        font_size=7.6,
    )

    add_note_box(
        doc,
        "未作为正证据放入的关键图",
        "F3 仍是 lightweight screening，不是最终 Koopman DNN 多 seed 预测；F12 当前标题显示 certificate ok=0.0% 且 min margin 为负，"
        "不能作为稳定性正证据；F13 是 available variants，不是真正逐模块 ablation；F11 的 phase-role 收益较弱，建议待 hairpin/S-curve 高曲率实验后再放主文。",
        WARN_FILL,
    )

    add_figure(
        doc,
        "fig05_main_trajectory_compare.png",
        "图 14-1  高通信退化和混合故障/噪声场景下的团队中心轨迹对比。TF14 保持完整路径跟踪，未出现轨迹发散。",
        width_in=6.4,
    )
    add_figure(
        doc,
        "fig06_tracking_error_statistics.png",
        "图 14-2  系统级横向/纵向误差与多场景 RMSE 统计。该图作为 TF14 控制效果优于 baseline 的主要误差证据。",
        width_in=6.4,
    )
    add_figure(
        doc,
        "fig08_fdi_identification.png",
        "图 14-3  在线 FDI 与执行器效率辨识。故障后残差、效率估计、置信度和模式识别形成完整诊断链。",
        width_in=6.2,
    )
    add_figure(
        doc,
        "fig09_ftc_switch_redistribution.png",
        "图 14-4  FTC 切换、车辆级控制指令和重分配补偿。该图证明故障后控制器不是被动跟踪，而是主动重构控制输入。",
        width_in=6.2,
    )
    add_figure(
        doc,
        "fig10_payload_force_connection_safety.png",
        "图 14-5  载荷受力和连接安全。TF14 的连接误差和连接裕度占用显著降低，但受力峰值存在主动补偿代价。",
        width_in=6.3,
    )
    add_figure(
        doc,
        "figS11_all_scenario_errors.png",
        "图 14-6  四类场景下的横向和纵向误差时间序列。该图作为主文误差统计之外的全场景补充证据。",
        width_in=6.3,
    )
    add_figure(
        doc,
        "figS15_control_saturation.png",
        "图 14-7  车辆级控制饱和检查。补偿后的转角和加速度仍处于约束边界内，支撑控制效果的可执行性。",
        width_in=6.3,
    )
    add_figure(
        doc,
        "figS19_tf13_error_match_runtime.png",
        "图 14-8  实时性、跟踪误差与完整路径成功率对比。TF14 main/phase-role 保持实时性，error-match/all-solve 提供精度-计算时间上界。",
        width_in=6.3,
    )

    add_para(doc, "15. 改进点与实验验证覆盖审查", style="Heading 1")
    add_para(
        doc,
        "下表检查 TF14 的每个主要改进是否已有对应实验验证。结论是：控制效果、FDI、FTC、连接安全和实时性已有可用证据；"
        "Koopman 网络优势、真实逐模块消融、Lyapunov 证书和 phase-role 的高曲率收益仍需要补实验。",
    )

    add_table(
        doc,
        ["改进点", "已有图/实验", "覆盖结论", "缺口与下一步"],
        [
            [
                "双线性 Koopman + 稳定投影",
                "F2、S01、F3",
                "部分覆盖。F2/S01 能证明数据和训练流程；F3 目前不能作为最终模型优势证据。",
                "重跑 E1：linear、bilinear、stable bilinear、online adaptation，各 >=5 seed，输出 one-step/multi-step/per-state 误差。",
            ],
            [
                "通信质量一致性、时延补偿和约束收紧",
                "F5、F6、S05、S11",
                "部分覆盖。高通信退化场景下控制仍稳定，S05 展示网络质量变化。",
                "还缺 no-comm-aware、delay-only、full TF14 的消融；需证明通信模块本身贡献，而不只是系统能跑。",
            ],
            [
                "在线 FDI 与执行器效率辨识",
                "F8、S06",
                "覆盖较好。F8 有检测、效率估计、置信度和模式识别；S06 有调参延迟与误报信息。",
                "补 S7：ax/steer/dual/severe/ramp/intermittent 的混淆矩阵和 ROC，证明分类鲁棒性。",
            ],
            [
                "切换 FTC 与控制重分配",
                "F9、S15",
                "覆盖较好。F9 展示切换和车辆级输入，S15 证明输入未越界。",
                "补 severe dual loss 和 safe degraded fallback 场景，证明安全降级分支真的有效。",
            ],
            [
                "载荷连接安全与协同运输物理量",
                "F10",
                "覆盖较好但需谨慎表述。连接误差明显降低，说明协同运输安全性提升。",
                "补所有场景的 Fx/Fy/Mz/RMS/peak/violation count；确认 baseline 与 TF14 使用相同扰动，避免公平性争议。",
            ],
            [
                "阶段-角色调度器",
                "F11、S12",
                "覆盖偏弱。当前 DLC 下收益不够突出，只能作为辅助机制图。",
                "补 hairpin、S-curve、narrow passage 高曲率实验；按 inner/outer、front/rear、fault/support 聚合误差与载荷力矩。",
            ],
            [
                "Lyapunov 型稳定性证书",
                "F12",
                "当前不能作为正证据。F12 标题显示 ok=0.0%、min margin 为负。",
                "需要检查 certificate_ok 统计口径、扰动界和 margin 定义；若理论是 practical/ISS 稳定，图中应展示 bounded margin 而不是 ok=0。",
            ],
            [
                "实时预算与 all-solve 精度-计算折中",
                "F14、S19",
                "覆盖较好。主实时 preset 和 all-solve/error-match 的计算代价差异清楚。",
                "补 p95/p99 step time、overrun ratio、vehicles_solved 和模块耗时分解。",
            ],
            [
                "完整逐模块消融",
                "F13",
                "未充分覆盖。F13 是 available variants，不是真正消融。",
                "必须重跑 E6：no bilinear、no stable projection、no comm、no FDI、no FTC、no phase-role、no realtime，每项只关一个模块。",
            ],
        ],
        [1.7, 1.5, 2.2, 3.0],
        font_size=7.4,
    )

    add_note_box(
        doc,
        "最终写作口径",
        "目前可以在中文稿中主张：TF14 在通信退化和故障场景下具有更好的轨迹保持、横向误差、在线故障检测、容错重分配、连接安全和实时可执行性。"
        "暂时不要主张：Koopman 网络已被最终多 seed 证明优于全部替代模型、phase-role 已显著改善所有场景、Lyapunov 证书已经实验正验证、F13 已完成完整消融。",
        WARN_FILL,
    )

    doc.save(str(OUT_DOCX))


if __name__ == "__main__":
    build_docx()
    print(OUT_DOCX)
